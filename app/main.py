import csv
import hmac
import json
import math
import os
import re
import secrets
import time
from contextlib import contextmanager
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urlparse

import httpx
import jwt
import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from fastapi import Body, FastAPI, HTTPException, Query, Request
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse, Response
from jwt import PyJWKClient
from psycopg.rows import dict_row
from psycopg.types.json import Json

ROOT = Path(__file__).resolve().parent.parent
PUBLIC_BASE = os.environ.get("PUBLIC_BASE", "https://localhost").rstrip("/")
OIDC_INTERNAL = os.environ.get("OIDC_INTERNAL", "http://keycloak:8080/auth/realms/dip").rstrip("/")
REDIRECT_URI = PUBLIC_BASE + "/oauth/callback"
CLIENT_ID = "dip-web"
MONEY = Decimal("0.01")
LATIN = re.compile(r"[A-Za-z]")

VIEWERS, EDITORS, AUDITORS, ADMINS = "DiP-Viewers", "DiP-Editors", "DiP-Auditors", "DiP-Admins"
ALL_GROUPS = {VIEWERS, EDITORS, AUDITORS, ADMINS}
EDIT_GROUPS = {EDITORS, ADMINS}
AUDIT_GROUPS = {EDITORS, AUDITORS, ADMINS}
ROLLBACK_GROUPS = {ADMINS}


def read_secret(name: str) -> str:
    for path in (Path("/run/secrets") / name, ROOT / "secrets" / name):
        if path.is_file():
            return path.read_text(encoding="utf-8").strip()
    raise SystemExit(f"missing secret: {name}")


SESSION_SECRET = read_secret("session_secret")
if len(SESSION_SECRET) < 32:
    raise SystemExit("session_secret must be at least 32 characters")
DB_PASSWORD = read_secret("postgres_password")
CLIENT_SECRET = read_secret("kc_client_secret")

SCHEMA = """
CREATE TABLE IF NOT EXISTS contracts (
  id text PRIMARY KEY,
  number text NOT NULL DEFAULT '',
  date text NOT NULL DEFAULT '',
  counterparty text NOT NULL DEFAULT '',
  subject text NOT NULL DEFAULT '',
  direction text NOT NULL DEFAULT '',
  kind text NOT NULL DEFAULT '',
  amount numeric(18,2) NOT NULL DEFAULT 0,
  vat_rate numeric(5,2) NOT NULL DEFAULT 0,
  project text NOT NULL DEFAULT '',
  responsible text NOT NULL DEFAULT '',
  start_date text NOT NULL DEFAULT '',
  end_date text NOT NULL DEFAULT '',
  status text NOT NULL DEFAULT 'active',
  comment text NOT NULL DEFAULT '',
  review_date text NOT NULL DEFAULT '',
  registration_date text NOT NULL DEFAULT '',
  document_card text NOT NULL DEFAULT '',
  funding_source text NOT NULL DEFAULT '',
  amount_without_vat numeric(18,2) NOT NULL DEFAULT 0,
  document_kind text NOT NULL DEFAULT 'contract',
  category text NOT NULL DEFAULT 'limit',
  other_expense boolean NOT NULL DEFAULT false,
  version integer NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS specifications (
  id text PRIMARY KEY,
  contract_id text NOT NULL REFERENCES contracts(id) ON DELETE CASCADE,
  number text NOT NULL DEFAULT '',
  date text NOT NULL DEFAULT '',
  subject text NOT NULL DEFAULT '',
  amount numeric(18,2) NOT NULL DEFAULT 0,
  vat_rate numeric(5,2) NOT NULL DEFAULT 0,
  status text NOT NULL DEFAULT 'active',
  version integer NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS payments (
  id text PRIMARY KEY,
  contract_id text NOT NULL REFERENCES contracts(id) ON DELETE CASCADE,
  number text NOT NULL DEFAULT '',
  date text NOT NULL DEFAULT '',
  amount numeric(18,2) NOT NULL DEFAULT 0,
  type text NOT NULL DEFAULT '',
  payment_kind text NOT NULL DEFAULT '',
  status text NOT NULL DEFAULT 'paid',
  comment text NOT NULL DEFAULT '',
  approval_date text NOT NULL DEFAULT '',
  payment_date text NOT NULL DEFAULT '',
  document_card text NOT NULL DEFAULT '',
  responsible text NOT NULL DEFAULT '',
  cfo text NOT NULL DEFAULT '',
  bdds text NOT NULL DEFAULT '',
  approval_amount numeric(18,2) NOT NULL DEFAULT 0,
  approval_vat_rate numeric(5,2) NOT NULL DEFAULT 0,
  specification_id text NOT NULL DEFAULT '',
  approval_amount_without_vat numeric(18,2) NOT NULL DEFAULT 0,
  approval_vat_amount numeric(18,2) NOT NULL DEFAULT 0,
  document_kind text NOT NULL DEFAULT '',
  subtype text NOT NULL DEFAULT '',
  approval_end_date text NOT NULL DEFAULT '',
  posted_date text NOT NULL DEFAULT '',
  cost_type text NOT NULL DEFAULT '',
  mvz text NOT NULL DEFAULT '',
  key_counterparty text NOT NULL DEFAULT '',
  project text NOT NULL DEFAULT '',
  version integer NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS sessions (
  id text PRIMARY KEY,
  username text NOT NULL,
  groups jsonb NOT NULL,
  csrf_token text NOT NULL,
  id_token text NOT NULL DEFAULT '',
  expires_at timestamptz NOT NULL
);
CREATE TABLE IF NOT EXISTS audit_log (
  id bigserial PRIMARY KEY,
  ts timestamptz NOT NULL DEFAULT now(),
  actor text NOT NULL,
  action text NOT NULL,
  entity_type text NOT NULL,
  entity_id text NOT NULL,
  label text NOT NULL,
  before_state jsonb,
  after_state jsonb
);
CREATE TABLE IF NOT EXISTS dictionaries (
  id text PRIMARY KEY,
  kind text NOT NULL,
  value text NOT NULL DEFAULT '',
  version integer NOT NULL DEFAULT 1
);
CREATE TABLE IF NOT EXISTS addendums (
  id text PRIMARY KEY,
  contract_id text NOT NULL REFERENCES contracts(id) ON DELETE CASCADE,
  number text NOT NULL DEFAULT '',
  document_card text NOT NULL DEFAULT '',
  approval_date text NOT NULL DEFAULT '',
  term_change text NOT NULL DEFAULT '',
  amount numeric(18,2) NOT NULL DEFAULT 0,
  vat_rate numeric(5,2) NOT NULL DEFAULT 0,
  advance_percent numeric(6,2),
  responsible text NOT NULL DEFAULT '',
  comment text NOT NULL DEFAULT '',
  created_at timestamptz NOT NULL DEFAULT now(),
  version integer NOT NULL DEFAULT 1
);
ALTER TABLE addendums ADD COLUMN IF NOT EXISTS end_date text NOT NULL DEFAULT '';
ALTER TABLE contracts ADD COLUMN IF NOT EXISTS parent_contract_id text NOT NULL DEFAULT '';
ALTER TABLE contracts ADD COLUMN IF NOT EXISTS funding_contract_id text NOT NULL DEFAULT '';
ALTER TABLE contracts ADD COLUMN IF NOT EXISTS advance_percent numeric(6,2);
ALTER TABLE payments ADD COLUMN IF NOT EXISTS counterparty text NOT NULL DEFAULT '';
ALTER TABLE payments ADD COLUMN IF NOT EXISTS funding_contract_id text NOT NULL DEFAULT '';
ALTER TABLE payments ALTER COLUMN contract_id DROP NOT NULL;
"""

DICTIONARY_KINDS = ("cfo", "bdds", "project")


def connect():
    return psycopg.connect(
        host=os.environ.get("DATABASE_HOST", "postgres"),
        dbname=os.environ.get("DATABASE_NAME", "dip"),
        user=os.environ.get("DATABASE_USER", "dip"),
        password=DB_PASSWORD,
        row_factory=dict_row,
    )


@contextmanager
def transaction():
    conn = connect()
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def optional_percent(value):
    raw = text(value).strip().replace("%", "").replace(",", ".")
    if not raw:
        return None
    number = parse_money(raw)
    if number is None:
        raise HTTPException(400, "Укажите процент числом")
    return number


def money(value) -> Decimal:
    number = parse_money(value)
    if number is None:
        raise HTTPException(400, "Укажите сумму числом")
    return number


def parse_money(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, Decimal):
        raw = value
    elif isinstance(value, int):
        raw = Decimal(value)
    elif isinstance(value, float):
        raw = Decimal(str(value))
    else:
        cleaned = str(value or "").strip().replace("\u00a0", " ").replace("\u202f", " ")
        for mark in ("₽", "руб.", "руб"):
            cleaned = cleaned.replace(mark, "")
        cleaned = cleaned.replace(" ", "")
        if not cleaned:
            return Decimal(0).quantize(MONEY)
        if cleaned.lower() in {"nan", "+nan", "-nan", "inf", "+inf", "-inf", "infinity", "+infinity", "-infinity", "snan"}:
            return None
        last_comma, last_dot = cleaned.rfind(","), cleaned.rfind(".")
        if last_comma >= 0 and last_dot >= 0:
            if last_comma > last_dot:
                cleaned = cleaned.replace(".", "").replace(",", ".")
            else:
                cleaned = cleaned.replace(",", "")
        elif last_comma >= 0:
            cleaned = cleaned.replace(",", ".")
        try:
            raw = Decimal(cleaned)
        except InvalidOperation:
            return None
    if not raw.is_finite():
        return None
    try:
        return raw.quantize(MONEY, rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return None


def text(value) -> str:
    return "" if value is None else str(value)


def as_bool(value) -> bool:
    return value is True or value in ("true", "on", "1", 1)


def without_vat(amount: Decimal, rate: Decimal) -> Decimal:
    if rate == 0:
        return amount
    return (amount / (1 + rate / Decimal(100))).quantize(MONEY, rounding=ROUND_HALF_UP)


def json_money(value):
    if isinstance(value, Decimal):
        return float(value)
    return value


def sign(value: str) -> str:
    mac = hmac.new(SESSION_SECRET.encode(), value.encode(), "sha256").hexdigest()
    return f"{value}.{mac}"


def unsign(value: str) -> str | None:
    if not value or "." not in value:
        return None
    raw, mac = value.rsplit(".", 1)
    expected = hmac.new(SESSION_SECRET.encode(), raw.encode(), "sha256").hexdigest()
    if len(mac) != len(expected) or not hmac.compare_digest(mac, expected):
        return None
    return raw


def cookie_flags(response, name: str, value: str, max_age: int):
    response.set_cookie(name, value, httponly=True, secure=True, samesite="lax", max_age=max_age, path="/")


def same_token(left: str, right: str) -> bool:
    if not left or not right or len(left) != len(right):
        return False
    return hmac.compare_digest(left, right)


def check_registry(number, card, responsible, comment):
    if responsible and re.search(r"\d", text(responsible)):
        raise HTTPException(400, "В ФИО не может быть цифр")
    if responsible and LATIN.search(text(responsible)):
        raise HTTPException(400, "Ввод данных в столбце Ответств. только на русском языке")
    if comment and LATIN.search(text(comment)):
        raise HTTPException(400, "Ввод данных в столбце Коммент. только на русском языке")


def audit(conn, actor: str, action: str, entity_type: str, entity_id: str, label: str, before, after):
    conn.execute(
        """INSERT INTO audit_log (actor, action, entity_type, entity_id, label, before_state, after_state)
           VALUES (%s, %s, %s, %s, %s, %s, %s)""",
        (actor, action, entity_type, entity_id, label, Json(before), Json(after)),
    )
    print(json.dumps({
        "event": "dip.audit",
        "actor": actor,
        "action": action,
        "entityType": entity_type,
        "entityId": entity_id,
        "label": label,
        "before": before,
        "after": after,
    }, ensure_ascii=False, default=str), flush=True)


def contract_from_row(row) -> dict:
    return {
        "id": row["id"],
        "number": row["number"],
        "date": row["date"],
        "counterparty": row["counterparty"],
        "subject": row["subject"],
        "direction": row["direction"],
        "kind": row["kind"],
        "amount": json_money(row["amount"]),
        "vatRate": json_money(row["vat_rate"]),
        "project": row["project"],
        "responsible": row["responsible"],
        "startDate": row["start_date"],
        "endDate": row["end_date"],
        "status": row["status"],
        "comment": row["comment"],
        "reviewDate": row["review_date"],
        "registrationDate": row["registration_date"],
        "documentCard": row["document_card"],
        "fundingSource": row["funding_source"],
        "amountWithoutVat": json_money(row["amount_without_vat"]),
        "documentKind": row["document_kind"],
        "category": row["category"],
        "otherExpense": row["other_expense"],
        "parentContractId": row.get("parent_contract_id") or "",
        "fundingContractId": row.get("funding_contract_id") or "",
        "advancePercent": None if row.get("advance_percent") is None else json_money(row["advance_percent"]),
        "version": row["version"],
    }


def dictionary_from_row(row) -> dict:
    return {"id": row["id"], "kind": row["kind"], "value": row["value"], "version": row["version"]}


def spec_from_row(row) -> dict:
    return {
        "id": row["id"],
        "contractId": row["contract_id"],
        "number": row["number"],
        "date": row["date"],
        "subject": row["subject"],
        "amount": json_money(row["amount"]),
        "vatRate": json_money(row["vat_rate"]),
        "status": row["status"],
        "version": row["version"],
    }


def addendum_from_row(row) -> dict:
    created = row.get("created_at")
    return {
        "id": row["id"],
        "contractId": row["contract_id"],
        "number": row["number"],
        "documentCard": row["document_card"],
        "approvalDate": row["approval_date"],
        "endDate": row["end_date"],
        "amount": json_money(row["amount"]),
        "vatRate": json_money(row["vat_rate"]),
        "advancePercent": None if row.get("advance_percent") is None else json_money(row["advance_percent"]),
        "responsible": row["responsible"],
        "comment": row["comment"],
        "createdAt": created.isoformat() if created else "",
        "version": row["version"],
    }


def payment_from_row(row) -> dict:
    return {
        "id": row["id"],
        "contractId": row["contract_id"] or "",
        "number": row["number"],
        "date": row["date"],
        "amount": json_money(row["amount"]),
        "type": row["type"],
        "paymentKind": row["payment_kind"],
        "status": row["status"],
        "comment": row["comment"],
        "approvalDate": row["approval_date"],
        "paymentDate": row["payment_date"],
        "documentCard": row["document_card"],
        "responsible": row["responsible"],
        "cfo": row["cfo"],
        "bdds": row["bdds"],
        "approvalAmount": json_money(row["approval_amount"]),
        "approvalVatRate": json_money(row["approval_vat_rate"]),
        "specificationId": row["specification_id"],
        "approvalAmountWithoutVat": json_money(row["approval_amount_without_vat"]),
        "approvalVatAmount": json_money(row["approval_vat_amount"]),
        "documentKind": row["document_kind"],
        "subtype": row["subtype"],
        "approvalEndDate": row["approval_end_date"],
        "postedDate": row["posted_date"],
        "costType": row["cost_type"],
        "mvz": row["mvz"],
        "keyCounterparty": row["key_counterparty"],
        "project": row["project"],
        "counterparty": row.get("counterparty") or "",
        "fundingContractId": row.get("funding_contract_id") or "",
        "version": row["version"],
    }


def load_contract(conn, entity_id: str):
    row = conn.execute("SELECT * FROM contracts WHERE id = %s", (entity_id,)).fetchone()
    return contract_from_row(row) if row else None


def load_spec(conn, entity_id: str):
    row = conn.execute("SELECT * FROM specifications WHERE id = %s", (entity_id,)).fetchone()
    return spec_from_row(row) if row else None


def load_payment(conn, entity_id: str):
    row = conn.execute("SELECT * FROM payments WHERE id = %s", (entity_id,)).fetchone()
    return payment_from_row(row) if row else None


def load_addendum(conn, entity_id: str):
    row = conn.execute("SELECT * FROM addendums WHERE id = %s", (entity_id,)).fetchone()
    return addendum_from_row(row) if row else None


def load_dictionary(conn, entity_id: str):
    row = conn.execute("SELECT * FROM dictionaries WHERE id = %s", (entity_id,)).fetchone()
    return dictionary_from_row(row) if row else None


DICTIONARY_LABELS = {"cfo": "ЦФО", "bdds": "Статья БДДС", "project": "Проект"}


def addendum_order_key(item):
    """Порядок ДС — по времени добавления: крайнее добавленное считается действующим."""
    return text(item.get("createdAt"))


def latest_addendum(items):
    """Действующее доп. соглашение — последнее добавленное."""
    if not items:
        return None
    return max(items, key=addendum_order_key)


def apply_addendums(contracts, addendums):
    """Подтягивает в договор условия действующего доп. соглашения; исходные значения сохраняются в base*."""
    by_contract: dict[str, list] = {}
    for item in addendums:
        by_contract.setdefault(item["contractId"], []).append(item)
    for contract in contracts:
        contract["baseAmount"] = contract["amount"]
        contract["baseVatRate"] = contract["vatRate"]
        contract["baseAmountWithoutVat"] = contract["amountWithoutVat"]
        contract["baseAdvancePercent"] = contract["advancePercent"]
        contract["baseEndDate"] = contract["endDate"]
        items = by_contract.get(contract["id"], [])
        contract["addendumCount"] = len(items)
        current = latest_addendum(items)
        contract["currentAddendumNumber"] = current["number"] if current else ""
        if not current:
            continue
        amount = Decimal(str(current["amount"]))
        rate = Decimal(str(current["vatRate"]))
        contract["amount"] = float(amount)
        contract["vatRate"] = float(rate)
        contract["amountWithoutVat"] = float(without_vat(amount, rate))
        # % авансирования берётся из последнего ДС, где он указан; если нигде не указан — из договора
        ordered = sorted(items, key=addendum_order_key)
        for item in reversed(ordered):
            if item.get("advancePercent") is not None:
                contract["advancePercent"] = item["advancePercent"]
                break
        # срок окончания — из последнего ДС, где он указан; иначе остаётся срок договора
        for item in reversed(ordered):
            if text(item.get("endDate")).strip():
                contract["endDate"] = item["endDate"]
                break


def decorate(contracts, specs, payments, addendums=()):
    apply_addendums(contracts, list(addendums))
    paid_by_contract: dict[str, Decimal] = {}
    for item in payments:
        if item["status"] == "paid" and item["contractId"]:
            paid_by_contract[item["contractId"]] = paid_by_contract.get(item["contractId"], Decimal(0)) + Decimal(str(item["amount"]))
    children: dict[str, list] = {}
    for contract in contracts:
        parent = contract.get("parentContractId") or ""
        if parent:
            children.setdefault(parent, []).append(contract)
    for contract in contracts:
        amount = Decimal(str(contract["amount"]))
        own_children = children.get(contract["id"], [])
        child_sum = sum((Decimal(str(item["amount"])) for item in own_children), Decimal(0))
        paid = paid_by_contract.get(contract["id"], Decimal(0))
        paid += sum((paid_by_contract.get(item["id"], Decimal(0)) for item in own_children), Decimal(0))
        contract["childCount"] = len(own_children)
        contract["childLimit"] = float(child_sum)
        contract["limitRemainder"] = None if contract["kind"] != "framework" else float(amount - child_sum)
        contract["payRemainder"] = float(amount - paid)
        contract["paidTotal"] = float(paid)
    for spec in specs:
        paid = sum(
            (Decimal(str(item["amount"])) for item in payments if item["specificationId"] == spec["id"] and item["status"] == "paid"),
            Decimal(0),
        )
        spec["payRemainder"] = float(Decimal(str(spec["amount"])) - paid)
    return contracts, specs, payments


def fetch_all(conn):
    contracts = [contract_from_row(row) for row in conn.execute("SELECT * FROM contracts ORDER BY number")]
    specs = [spec_from_row(row) for row in conn.execute("SELECT * FROM specifications ORDER BY number")]
    payments = [payment_from_row(row) for row in conn.execute("SELECT * FROM payments ORDER BY date DESC")]
    addendums = [addendum_from_row(row) for row in conn.execute("SELECT * FROM addendums ORDER BY created_at DESC")]
    contracts, specs, payments = decorate(contracts, specs, payments, addendums)
    return contracts, specs, payments, addendums


def normalize_contract(payload: dict, existing: dict | None) -> dict:
    check_registry(payload.get("number"), payload.get("documentCard"), payload.get("responsible"), payload.get("comment"))
    document_kind = text(payload.get("documentKind") or "contract")
    direction = text(payload.get("direction"))
    kind = text(payload.get("kind"))
    category = text(payload.get("category") or "limit")
    start_date = text(payload.get("startDate"))
    end_date = text(payload.get("endDate"))
    parent_id = text(payload.get("parentContractId")).strip()
    if category == "addendum":
        raise HTTPException(400, "Доп. соглашения ведутся на отдельном листе «Доп. соглашения»")
    if not text(payload.get("subject")).strip():
        raise HTTPException(400, "Укажите предмет договора")
    if category == "limit":
        parent_id = ""
    elif not parent_id:
        raise HTTPException(400, "Для спецификации или заявки укажите договор-основание")
    if parent_id:
        document_kind = "contract"
    if document_kind in ("invoice", "none"):
        direction, kind, category, start_date, end_date = "", "", "limit", "", ""
    elif kind == "fixed" and category in ("specification", "request") and not parent_id:
        category = "limit"
    amount = money(payload.get("amount"))
    rate = money(payload.get("vatRate"))
    review = text(payload.get("reviewDate") or payload.get("date"))
    funding_id = text(payload.get("fundingContractId")).strip()
    return {
        "id": text(payload.get("id") or (existing or {}).get("id") or "contract-" + secrets.token_hex(8)),
        "number": text(payload.get("number")).strip(),
        "date": review,
        "counterparty": text(payload.get("counterparty")),
        "subject": text(payload.get("subject")),
        "direction": direction,
        "kind": kind,
        "amount": amount,
        "vatRate": rate,
        "project": text(payload.get("project")),
        "responsible": text(payload.get("responsible")),
        "startDate": start_date,
        "endDate": end_date,
        "status": text(payload.get("status") or (existing or {}).get("status") or "active"),
        "comment": text(payload.get("comment")),
        "reviewDate": review,
        "registrationDate": text(payload.get("registrationDate") or review),
        "documentCard": text(payload.get("documentCard")).strip(),
        "fundingSource": text(payload.get("fundingSource")),
        "amountWithoutVat": without_vat(amount, rate),
        "documentKind": document_kind,
        "category": category or "limit",
        "otherExpense": direction == "expense" and as_bool(payload.get("otherExpense")),
        "parentContractId": parent_id,
        "fundingContractId": funding_id,
        "advancePercent": optional_percent(payload.get("advancePercent")),
        "version": int((existing or {}).get("version") or 1),
    }


def normalize_spec(payload: dict, existing: dict | None) -> dict:
    amount = money(payload.get("amount"))
    rate = money(payload.get("vatRate"))
    return {
        "id": text(payload.get("id") or (existing or {}).get("id") or "spec-" + secrets.token_hex(8)),
        "contractId": text(payload.get("contractId")),
        "number": text(payload.get("number")),
        "date": text(payload.get("date")),
        "subject": text(payload.get("subject")),
        "amount": amount,
        "vatRate": rate,
        "status": text(payload.get("status") or (existing or {}).get("status") or "active"),
        "version": int((existing or {}).get("version") or 1),
    }


def normalize_payment(payload: dict, existing: dict | None) -> dict:
    check_registry(payload.get("number"), payload.get("documentCard"), payload.get("responsible"), payload.get("comment"))
    amount = money(payload.get("amount"))
    approval = money(payload.get("approvalAmount") if payload.get("approvalAmount") not in (None, "") else amount)
    rate = money(payload.get("approvalVatRate"))
    net = without_vat(approval, rate)
    payment_date = text(payload.get("paymentDate")).strip() if "paymentDate" in payload else text(payload.get("date")).strip()
    approval_date = text(payload.get("approvalDate") or payment_date).strip()
    subtype = text(payload.get("subtype"))
    document_kind = text(payload.get("documentKind") or "contract")
    contract_id = text(payload.get("contractId")).strip()
    number = text(payload.get("number")).strip()
    if document_kind != "none" and not number:
        raise HTTPException(400, "Укажите номер документа")
    if document_kind == "contract" and not contract_id:
        raise HTTPException(400, "Для счета по договору выберите договор")
    if document_kind != "contract":
        contract_id = ""
    if not text(payload.get("bdds")).strip():
        raise HTTPException(400, "Укажите статью БДДС")
    if not text(payload.get("cfo")).strip():
        raise HTTPException(400, "Укажите ЦФО")
    return {
        "id": text(payload.get("id") or (existing or {}).get("id") or "payment-" + secrets.token_hex(8)),
        "contractId": contract_id,
        "number": number,
        "date": payment_date,
        "amount": amount,
        "type": text(payload.get("type")),
        "paymentKind": text(payload.get("paymentKind")),
        "status": text(payload.get("status") or (existing or {}).get("status") or "paid"),
        "comment": text(payload.get("comment")),
        "approvalDate": approval_date,
        "paymentDate": payment_date,
        "documentCard": text(payload.get("documentCard")).strip(),
        "responsible": text(payload.get("responsible")),
        "cfo": text(payload.get("cfo")),
        "bdds": text(payload.get("bdds")),
        "approvalAmount": approval,
        "approvalVatRate": rate,
        "specificationId": "" if subtype != "specification" else text(payload.get("specificationId")),
        "approvalAmountWithoutVat": net,
        "approvalVatAmount": (approval - net).quantize(MONEY, rounding=ROUND_HALF_UP),
        "documentKind": document_kind,
        "subtype": subtype,
        "approvalEndDate": text(payload.get("approvalEndDate")).strip() if "approvalEndDate" in payload else text((existing or {}).get("approvalEndDate") or ""),
        "postedDate": text(payload.get("postedDate")).strip() if "postedDate" in payload else text((existing or {}).get("postedDate") or ""),
        "costType": text(payload.get("costType")) if "costType" in payload else text((existing or {}).get("costType") or payload.get("bdds")),
        "mvz": text(payload.get("mvz")) if "mvz" in payload else text((existing or {}).get("mvz") or payload.get("cfo")),
        "keyCounterparty": text(payload.get("keyCounterparty")) if "keyCounterparty" in payload else text((existing or {}).get("keyCounterparty") or payload.get("counterparty")),
        "project": text(payload.get("project")),
        "counterparty": text(payload.get("counterparty")).strip(),
        "fundingContractId": text(payload.get("fundingContractId")).strip(),
        "version": int((existing or {}).get("version") or 1),
    }


def normalize_addendum(payload: dict, existing: dict | None) -> dict:
    check_registry(payload.get("number"), payload.get("documentCard"), payload.get("responsible"), payload.get("comment"))
    contract_id = text(payload.get("contractId")).strip()
    if not contract_id:
        raise HTTPException(400, "Выберите договор, к которому относится доп. соглашение")
    number = text(payload.get("number")).strip()
    if not number:
        raise HTTPException(400, "Укажите номер доп. соглашения")
    # В доп. соглашении обязательны все поля
    document_card = text(payload.get("documentCard")).strip()
    if not document_card:
        raise HTTPException(400, "Укажите карту документа доп. соглашения")
    approval_date = text(payload.get("approvalDate")).strip()[:10]
    try:
        datetime.strptime(approval_date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(400, "Укажите дату согласования доп. соглашения")
    amount = money(payload.get("amount"))
    rate = money(payload.get("vatRate"))
    end_date = text(payload.get("endDate")).strip()[:10]
    try:
        datetime.strptime(end_date, "%Y-%m-%d")
    except ValueError:
        raise HTTPException(400, "Укажите срок окончания по доп. соглашению (дата)")
    advance_percent = optional_percent(payload.get("advancePercent"))
    if advance_percent is None:
        raise HTTPException(400, "Укажите % авансирования по доп. соглашению")
    responsible = text(payload.get("responsible")).strip()
    if not responsible:
        raise HTTPException(400, "Укажите ответственного по доп. соглашению")
    comment = text(payload.get("comment")).strip()
    if not comment:
        raise HTTPException(400, "Укажите комментарий к доп. соглашению")
    return {
        "id": text(payload.get("id") or (existing or {}).get("id") or "addendum-" + secrets.token_hex(8)),
        "contractId": contract_id,
        "number": number,
        "documentCard": document_card,
        "approvalDate": approval_date,
        "endDate": end_date,
        "amount": amount,
        "vatRate": rate,
        "advancePercent": advance_percent,
        "responsible": responsible,
        "comment": comment,
        "createdAt": text((existing or {}).get("createdAt") or payload.get("createdAt") or ""),
        "version": int((existing or {}).get("version") or 1),
    }


def normalize_dictionary(payload: dict, existing: dict | None) -> dict:
    kind = text(payload.get("kind") or (existing or {}).get("kind")).strip()
    if kind not in DICTIONARY_KINDS:
        raise HTTPException(400, "Неизвестный справочник")
    value = text(payload.get("value")).strip()
    if not value:
        raise HTTPException(400, "Укажите значение")
    return {
        "id": text(payload.get("id") or (existing or {}).get("id") or "dict-" + secrets.token_hex(8)),
        "kind": kind,
        "value": value,
        "version": int((existing or {}).get("version") or 1),
    }


def public_copy(entity: dict) -> dict:
    copied = {}
    for key, value in entity.items():
        copied[key] = float(value) if isinstance(value, Decimal) else value
    return copied


CONTRACT_SQL = """
id, number, date, counterparty, subject, direction, kind, amount, vat_rate, project, responsible,
start_date, end_date, status, comment, review_date, registration_date, document_card, funding_source,
amount_without_vat, document_kind, category, other_expense, parent_contract_id, funding_contract_id, advance_percent, version
"""


def contract_params(entity: dict, version: int):
    return (
        entity["id"], entity["number"], entity["date"], entity["counterparty"], entity["subject"],
        entity["direction"], entity["kind"], entity["amount"], entity["vatRate"], entity["project"],
        entity["responsible"], entity["startDate"], entity["endDate"], entity["status"], entity["comment"],
        entity["reviewDate"], entity["registrationDate"], entity["documentCard"], entity["fundingSource"],
        entity["amountWithoutVat"], entity["documentKind"], entity["category"], entity["otherExpense"],
        entity.get("parentContractId") or "", entity.get("fundingContractId") or "", entity.get("advancePercent"), version,
    )


def upsert_contract(conn, entity: dict, version: int):
    parent_id = entity.get("parentContractId") or ""
    if parent_id:
        if parent_id == entity["id"]:
            raise HTTPException(400, "Договор не может быть основанием для самого себя")
        if conn.execute("SELECT 1 FROM contracts WHERE id = %s", (parent_id,)).fetchone() is None:
            raise HTTPException(400, "Договор-основание не найден")
    assignments = """
    number=%s, date=%s, counterparty=%s, subject=%s, direction=%s, kind=%s, amount=%s, vat_rate=%s,
    project=%s, responsible=%s, start_date=%s, end_date=%s, status=%s, comment=%s, review_date=%s,
    registration_date=%s, document_card=%s, funding_source=%s, amount_without_vat=%s, document_kind=%s,
    category=%s, other_expense=%s, parent_contract_id=%s, funding_contract_id=%s, advance_percent=%s, version=%s
    """
    params = contract_params(entity, version)[1:] + (entity["id"],)
    conn.execute(f"UPDATE contracts SET {assignments} WHERE id = %s", params)
    if conn.execute("SELECT 1 FROM contracts WHERE id = %s", (entity["id"],)).fetchone() is None:
        conn.execute(f"INSERT INTO contracts ({CONTRACT_SQL}) VALUES ({', '.join(['%s'] * 27)})", contract_params(entity, version))


def upsert_dictionary(conn, entity: dict, version: int):
    updated = conn.execute(
        "UPDATE dictionaries SET kind=%s, value=%s, version=%s WHERE id=%s",
        (entity["kind"], entity["value"], version, entity["id"]),
    )
    if updated.rowcount == 0:
        conn.execute(
            "INSERT INTO dictionaries (id, kind, value, version) VALUES (%s, %s, %s, %s)",
            (entity["id"], entity["kind"], entity["value"], version),
        )


def upsert_spec(conn, entity: dict, version: int):
    if conn.execute("SELECT 1 FROM contracts WHERE id = %s", (entity["contractId"],)).fetchone() is None:
        raise HTTPException(400, "Договор не найден")
    updated = conn.execute(
        """UPDATE specifications SET contract_id=%s, number=%s, date=%s, subject=%s, amount=%s, vat_rate=%s, status=%s, version=%s
           WHERE id=%s""",
        (entity["contractId"], entity["number"], entity["date"], entity["subject"], entity["amount"], entity["vatRate"], entity["status"], version, entity["id"]),
    )
    if updated.rowcount == 0:
        conn.execute(
            """INSERT INTO specifications (id, contract_id, number, date, subject, amount, vat_rate, status, version)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)""",
            (entity["id"], entity["contractId"], entity["number"], entity["date"], entity["subject"], entity["amount"], entity["vatRate"], entity["status"], version),
        )


def upsert_addendum(conn, entity: dict, version: int):
    parent = conn.execute("SELECT parent_contract_id, document_kind FROM contracts WHERE id = %s", (entity["contractId"],)).fetchone()
    if parent is None:
        raise HTTPException(400, "Договор не найден")
    if parent["parent_contract_id"]:
        raise HTTPException(400, "Доп. соглашение оформляется к договору, а не к спецификации")
    values = (
        entity["contractId"], entity["number"], entity["documentCard"], entity["approvalDate"], entity.get("endDate", ""),
        entity["amount"], entity["vatRate"], entity.get("advancePercent"), entity["responsible"], entity["comment"], version, entity["id"],
    )
    updated = conn.execute(
        """UPDATE addendums SET contract_id=%s, number=%s, document_card=%s, approval_date=%s, end_date=%s,
           amount=%s, vat_rate=%s, advance_percent=%s, responsible=%s, comment=%s, version=%s WHERE id=%s""",
        values,
    )
    if updated.rowcount == 0:
        created = entity.get("createdAt") or None
        conn.execute(
            """INSERT INTO addendums (id, contract_id, number, document_card, approval_date, end_date, amount, vat_rate,
                 advance_percent, responsible, comment, version, created_at)
               VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, COALESCE(%s::timestamptz, now()))""",
            (entity["id"], *values[:-1], created),
        )


def upsert_payment(conn, entity: dict, version: int):
    contract_id = entity.get("contractId") or ""
    if contract_id and conn.execute("SELECT 1 FROM contracts WHERE id = %s", (contract_id,)).fetchone() is None:
        raise HTTPException(400, "Договор не найден")
    values = (
        contract_id or None, entity["number"], entity["date"], entity["amount"], entity["type"], entity["paymentKind"],
        entity["status"], entity["comment"], entity["approvalDate"], entity["paymentDate"], entity["documentCard"],
        entity["responsible"], entity["cfo"], entity["bdds"], entity["approvalAmount"], entity["approvalVatRate"],
        entity["specificationId"], entity["approvalAmountWithoutVat"], entity["approvalVatAmount"], entity["documentKind"],
        entity["subtype"], entity["approvalEndDate"], entity["postedDate"], entity["costType"], entity["mvz"],
        entity["keyCounterparty"], entity["project"], entity.get("counterparty") or "", entity.get("fundingContractId") or "",
        version, entity["id"],
    )
    updated = conn.execute(
        """UPDATE payments SET contract_id=%s, number=%s, date=%s, amount=%s, type=%s, payment_kind=%s, status=%s,
           comment=%s, approval_date=%s, payment_date=%s, document_card=%s, responsible=%s, cfo=%s, bdds=%s,
           approval_amount=%s, approval_vat_rate=%s, specification_id=%s, approval_amount_without_vat=%s,
           approval_vat_amount=%s, document_kind=%s, subtype=%s, approval_end_date=%s, posted_date=%s, cost_type=%s,
           mvz=%s, key_counterparty=%s, project=%s, counterparty=%s, funding_contract_id=%s, version=%s WHERE id=%s""",
        values,
    )
    if updated.rowcount == 0:
        conn.execute(
            """INSERT INTO payments (
                 id, contract_id, number, date, amount, type, payment_kind, status, comment, approval_date, payment_date,
                 document_card, responsible, cfo, bdds, approval_amount, approval_vat_rate, specification_id,
                 approval_amount_without_vat, approval_vat_amount, document_kind, subtype, approval_end_date, posted_date,
                 cost_type, mvz, key_counterparty, project, counterparty, funding_contract_id, version
               ) VALUES (
                 %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s
               )""",
            (entity["id"], *values[:-1]),
        )


def contract_bundle(conn, entity_id: str):
    contract = load_contract(conn, entity_id)
    if not contract:
        return None
    specs = [spec_from_row(row) for row in conn.execute("SELECT * FROM specifications WHERE contract_id = %s", (entity_id,))]
    children = [contract_from_row(row) for row in conn.execute("SELECT * FROM contracts WHERE parent_contract_id = %s", (entity_id,))]
    ids = [entity_id] + [child["id"] for child in children]
    payments = [payment_from_row(row) for row in conn.execute("SELECT * FROM payments WHERE contract_id = ANY(%s)", (ids,))]
    addendums = [addendum_from_row(row) for row in conn.execute("SELECT * FROM addendums WHERE contract_id = %s", (entity_id,))]
    return {"entity": contract, "specifications": specs, "children": children, "payments": payments, "addendums": addendums}


def apply_saved(conn, kind: str, entity: dict, version: int):
    entity = dict(entity)
    entity["version"] = version
    if kind == "contract":
        upsert_contract(conn, entity, version)
    elif kind == "specification":
        upsert_spec(conn, entity, version)
    elif kind == "dictionary":
        upsert_dictionary(conn, entity, version)
    elif kind == "addendum":
        upsert_addendum(conn, entity, version)
    else:
        upsert_payment(conn, entity, version)
    return public_copy(entity)


def migrate_addendums(conn):
    """Дочерние договоры категории «Доп. соглашение» переносятся в таблицу addendums, платежи — на договор-основание."""
    rows = conn.execute("SELECT * FROM contracts WHERE category = 'addendum'").fetchall()
    for row in rows:
        item = contract_from_row(row)
        parent_id = item.get("parentContractId") or ""
        parent = load_contract(conn, parent_id) if parent_id else None
        if not parent:
            # без основания это обычный договор: категория была указана ошибочно
            conn.execute("UPDATE contracts SET category = 'limit', parent_contract_id = '' WHERE id = %s", (item["id"],))
            continue
        if load_addendum(conn, item["id"]) is None:
            upsert_addendum(conn, {
                "id": item["id"],
                "contractId": parent_id,
                "number": item["number"],
                "documentCard": item["documentCard"],
                "approvalDate": item["reviewDate"] or item["date"],
                "endDate": "",
                "amount": Decimal(str(item["amount"])),
                "vatRate": Decimal(str(item["vatRate"])),
                "advancePercent": item.get("advancePercent"),
                "responsible": item["responsible"],
                "comment": item["comment"],
                "createdAt": "",
            }, 1)
        conn.execute("UPDATE payments SET contract_id = %s WHERE contract_id = %s", (parent_id, item["id"]))
        conn.execute("UPDATE contracts SET funding_contract_id = '' WHERE funding_contract_id = %s", (item["id"],))
        conn.execute("UPDATE payments SET funding_contract_id = '' WHERE funding_contract_id = %s", (item["id"],))
        conn.execute("DELETE FROM contracts WHERE id = %s", (item["id"],))


def migrate_addendum_terms(conn):
    """Текстовое «Изменение сроков» переносится в дату окончания, если текст разбирается как дата."""
    rows = conn.execute("SELECT id, term_change FROM addendums WHERE end_date = '' AND term_change <> ''").fetchall()
    for row in rows:
        raw = text(row["term_change"]).strip()
        candidate = raw[:10] if re.match(r"\d{4}-\d{2}-\d{2}", raw) else raw
        for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
            try:
                parsed = datetime.strptime(candidate, pattern).date().isoformat()
            except ValueError:
                continue
            conn.execute("UPDATE addendums SET end_date = %s, term_change = '' WHERE id = %s", (parsed, row["id"]))
            break


def migrate_specifications(conn):
    """Переносит строки таблицы specifications в дочерние договоры и перепривязывает платежи."""
    rows = conn.execute("SELECT * FROM specifications").fetchall()
    if not rows:
        return
    for row in rows:
        spec = spec_from_row(row)
        parent = load_contract(conn, spec["contractId"])
        if not parent:
            conn.execute("DELETE FROM specifications WHERE id = %s", (spec["id"],))
            continue
        amount = Decimal(str(spec["amount"]))
        rate = Decimal(str(spec["vatRate"]))
        if load_contract(conn, spec["id"]) is None:
            upsert_contract(conn, {
                "id": spec["id"],
                "number": spec["number"],
                "date": spec["date"] or parent["date"],
                "counterparty": parent["counterparty"],
                "subject": spec["subject"],
                "direction": parent["direction"],
                "kind": "fixed",
                "amount": amount,
                "vatRate": rate,
                "project": parent["project"],
                "responsible": parent["responsible"],
                "startDate": "",
                "endDate": "",
                "status": spec["status"] or "active",
                "comment": "",
                "reviewDate": spec["date"] or parent["date"],
                "registrationDate": spec["date"] or parent["date"],
                "documentCard": "",
                "fundingSource": parent["fundingSource"],
                "amountWithoutVat": without_vat(amount, rate),
                "documentKind": "contract",
                "category": "specification",
                "otherExpense": False,
                "parentContractId": parent["id"],
                "fundingContractId": parent.get("fundingContractId") or "",
            }, 1)
        conn.execute(
            "UPDATE payments SET contract_id = %s, specification_id = '' WHERE specification_id = %s",
            (spec["id"], spec["id"]),
        )
        conn.execute("DELETE FROM specifications WHERE id = %s", (spec["id"],))


def seed_dictionaries(conn):
    if conn.execute("SELECT 1 FROM dictionaries LIMIT 1").fetchone():
        return
    sources = {
        "cfo": "SELECT DISTINCT cfo AS value FROM payments WHERE cfo <> ''",
        "bdds": "SELECT DISTINCT bdds AS value FROM payments WHERE bdds <> ''",
        "project": "SELECT DISTINCT project AS value FROM contracts WHERE project <> ''",
    }
    for kind, sql in sources.items():
        for row in conn.execute(sql).fetchall():
            value = text(row["value"]).strip()
            if value:
                upsert_dictionary(conn, {"id": "dict-" + secrets.token_hex(8), "kind": kind, "value": value}, 1)


def seed(conn):
    if conn.execute("SELECT 1 FROM contracts LIMIT 1").fetchone():
        return
    path = ROOT / "seed" / "data.json"
    if not path.is_file():
        return
    payload = json.loads(path.read_text(encoding="utf-8"))
    for raw in payload.get("contracts") or []:
        amount = money(raw.get("amount"))
        rate = money(raw.get("vatRate"))
        review = text(raw.get("reviewDate") or raw.get("date"))
        upsert_contract(conn, {
            "id": text(raw.get("id")),
            "number": text(raw.get("number")),
            "date": text(raw.get("date") or review),
            "counterparty": text(raw.get("counterparty")),
            "subject": text(raw.get("subject")),
            "direction": text(raw.get("direction")),
            "kind": text(raw.get("kind")),
            "amount": amount,
            "vatRate": rate,
            "project": text(raw.get("project")),
            "responsible": text(raw.get("responsible")),
            "startDate": text(raw.get("startDate")),
            "endDate": text(raw.get("endDate")),
            "status": text(raw.get("status") or "active"),
            "comment": text(raw.get("comment")),
            "reviewDate": review,
            "registrationDate": text(raw.get("registrationDate") or review),
            "documentCard": text(raw.get("documentCard")),
            "fundingSource": text(raw.get("fundingSource")),
            "amountWithoutVat": money(raw["amountWithoutVat"]) if raw.get("amountWithoutVat") not in (None, "") else without_vat(amount, rate),
            "documentKind": text(raw.get("documentKind") or "contract"),
            "category": text(raw.get("category") or "limit"),
            "otherExpense": as_bool(raw.get("otherExpense")),
            "parentContractId": text(raw.get("parentContractId")),
            "fundingContractId": text(raw.get("fundingContractId")),
            "advancePercent": optional_percent(raw.get("advancePercent")),
        }, 1)
    for raw in payload.get("specifications") or []:
        upsert_spec(conn, {
            "id": text(raw.get("id")),
            "contractId": text(raw.get("contractId")),
            "number": text(raw.get("number")),
            "date": text(raw.get("date")),
            "subject": text(raw.get("subject")),
            "amount": money(raw.get("amount")),
            "vatRate": money(raw.get("vatRate")),
            "status": text(raw.get("status") or "active"),
        }, 1)
    for raw in payload.get("payments") or []:
        amount = money(raw.get("amount"))
        approval = money(raw.get("approvalAmount") if raw.get("approvalAmount") not in (None, "") else amount)
        rate = money(raw.get("approvalVatRate"))
        net = money(raw["approvalAmountWithoutVat"]) if raw.get("approvalAmountWithoutVat") not in (None, "") else without_vat(approval, rate)
        vat_amount = money(raw["approvalVatAmount"]) if raw.get("approvalVatAmount") not in (None, "") else (approval - net)
        payment_date = text(raw.get("paymentDate") or raw.get("date"))
        upsert_payment(conn, {
            "id": text(raw.get("id")),
            "contractId": text(raw.get("contractId")),
            "number": text(raw.get("number")),
            "date": text(raw.get("date") or payment_date),
            "amount": amount,
            "type": text(raw.get("type")),
            "paymentKind": text(raw.get("paymentKind")),
            "status": text(raw.get("status") or "paid"),
            "comment": text(raw.get("comment")),
            "approvalDate": text(raw.get("approvalDate") or payment_date),
            "paymentDate": payment_date,
            "documentCard": text(raw.get("documentCard")),
            "responsible": text(raw.get("responsible")),
            "cfo": text(raw.get("cfo")),
            "bdds": text(raw.get("bdds")),
            "approvalAmount": approval,
            "approvalVatRate": rate,
            "specificationId": text(raw.get("specificationId")),
            "approvalAmountWithoutVat": net,
            "approvalVatAmount": vat_amount,
            "documentKind": text(raw.get("documentKind")),
            "subtype": text(raw.get("subtype")),
            "approvalEndDate": text(raw.get("approvalEndDate")),
            "postedDate": text(raw.get("postedDate")),
            "costType": text(raw.get("costType")),
            "mvz": text(raw.get("mvz")),
            "keyCounterparty": text(raw.get("keyCounterparty")),
            "project": text(raw.get("project")),
            "counterparty": text(raw.get("counterparty")),
            "fundingContractId": text(raw.get("fundingContractId")),
        }, 1)
    for raw in payload.get("dictionaries") or []:
        if text(raw.get("kind")) in DICTIONARY_KINDS and text(raw.get("value")).strip():
            upsert_dictionary(conn, {
                "id": text(raw.get("id") or "dict-" + secrets.token_hex(8)),
                "kind": text(raw.get("kind")),
                "value": text(raw.get("value")).strip(),
            }, 1)


def prepare_database():
    last_error = None
    for _ in range(30):
        try:
            with transaction() as conn:
                for statement in SCHEMA.split(";"):
                    sql = statement.strip()
                    if sql:
                        conn.execute(sql)
                seed(conn)
                migrate_specifications(conn)
                migrate_addendums(conn)
                migrate_addendum_terms(conn)
                conn.execute("UPDATE payments SET counterparty = c.counterparty FROM contracts c WHERE payments.contract_id = c.id AND payments.counterparty = ''")
                seed_dictionaries(conn)
            return
        except Exception as exc:
            last_error = exc
            print(f"database wait: {exc}", flush=True)
            time.sleep(2)
    raise SystemExit(f"database is not ready: {last_error}")


app = FastAPI(title="DiP registry")
_jwks = None


@app.on_event("startup")
def startup():
    prepare_database()


def jwks() -> PyJWKClient:
    global _jwks
    if _jwks is None:
        _jwks = PyJWKClient(OIDC_INTERNAL + "/protocol/openid-connect/certs")
    return _jwks


def session_row(conn, request: Request):
    raw = unsign(request.cookies.get("dip_session", ""))
    if not raw:
        return None
    return conn.execute("SELECT * FROM sessions WHERE id = %s AND expires_at > now()", (raw,)).fetchone()


def open_session(request: Request, mutating: bool):
    with transaction() as conn:
        row = session_row(conn, request)
    if not row:
        raise HTTPException(401, "Требуется вход")
    if mutating and not same_token(request.headers.get("x-csrf-token", ""), row["csrf_token"]):
        raise HTTPException(403, "CSRF")
    return row


def guard(request: Request, allowed: set[str]):
    row = open_session(request, request.method in {"POST", "PUT", "PATCH", "DELETE"})
    groups = set(row["groups"] or [])
    if not groups & ALL_GROUPS:
        raise HTTPException(403, "Нет группы доступа")
    if not groups & allowed:
        raise HTTPException(403, "Недостаточно прав")
    return row


def me_payload(row) -> dict:
    groups = list(row["groups"] or [])
    present = set(groups)
    return {
        "username": row["username"],
        "groups": groups,
        "allowed": bool(present & ALL_GROUPS),
        "canEdit": bool(present & EDIT_GROUPS),
        "canAudit": bool(present & AUDIT_GROUPS),
        "canRollback": bool(present & ROLLBACK_GROUPS),
        "csrf": row["csrf_token"],
    }


@app.get("/health")
def health():
    return {"status": "ok"}


@app.get("/ready")
def ready():
    try:
        with transaction() as conn:
            conn.execute("SELECT 1")
    except Exception:
        return JSONResponse({"status": "down"}, status_code=503)
    return {"status": "ok"}


def kc_headers(request: Request) -> dict:
    cookie = "; ".join(
        f"{name}={value}" for name, value in request.cookies.items() if name not in {"dip_session", "dip_oauth"}
    )
    headers = {
        "X-Forwarded-Proto": "https",
        "X-Forwarded-Host": "localhost",
        "X-Forwarded-Port": "443",
        "Accept": "text/html,application/xhtml+xml",
    }
    if cookie:
        headers["Cookie"] = cookie
    return headers


def to_internal(url: str) -> str:
    if url.startswith(PUBLIC_BASE + "/auth"):
        return "http://keycloak:8080" + url[len(PUBLIC_BASE):]
    if url.startswith("/auth"):
        return "http://keycloak:8080" + url
    return url


def public_cookie(raw: str) -> str:
    pieces = [part.strip() for part in raw.split(";")]
    kept = [pieces[0]]
    has_path = False
    for part in pieces[1:]:
        lower = part.lower()
        if lower.startswith("path="):
            kept.append("Path=/")
            has_path = True
        elif lower.startswith("domain="):
            continue
        else:
            kept.append(part)
    if not has_path:
        kept.append("Path=/")
    return "; ".join(kept)


def attach_cookies(response, raw_cookies: list[str]):
    for raw in raw_cookies:
        response.headers.append("set-cookie", public_cookie(raw))


def rewrite_login_page(html: str) -> tuple[str, str]:
    match = re.search(r'action="([^"]+)"', html)
    action = ""
    if match:
        action = match.group(1).replace("&amp;", "&")
        html = html[:match.start(1)] + "/login" + html[match.end(1):]
    return html, action


def login_page(request: Request, saved: dict) -> HTMLResponse:
    query = urlencode({
        "client_id": CLIENT_ID,
        "response_type": "code",
        "scope": "openid",
        "redirect_uri": REDIRECT_URI,
        "state": saved["state"],
        "nonce": saved["nonce"],
    })
    headers = kc_headers(request)
    cookies: list[str] = []
    url = OIDC_INTERNAL + "/protocol/openid-connect/auth?" + query
    with httpx.Client(timeout=20, follow_redirects=False) as client:
        upstream = client.get(url, headers=headers)
        cookies.extend(upstream.headers.get_list("set-cookie"))
        for _ in range(6):
            if upstream.status_code not in {301, 302, 303, 307, 308}:
                break
            location = upstream.headers.get("location", "")
            if "/oauth/callback" in location:
                parsed = urlparse(location)
                params = parse_qs(parsed.query)
                return complete_login(saved, params.get("code", [""])[0], params.get("state", [""])[0], cookies)
            headers["Cookie"] = merge_cookie(headers.get("Cookie", ""), cookies)
            upstream = client.get(to_internal(location), headers=headers)
            cookies.extend(upstream.headers.get_list("set-cookie"))
    if upstream.status_code != 200:
        raise HTTPException(502, "Страница входа недоступна")
    html, action = rewrite_login_page(upstream.text)
    if action:
        saved["action"] = action
    response = HTMLResponse(html)
    cookie_flags(response, "dip_oauth", sign(json.dumps(saved)), 600)
    attach_cookies(response, cookies)
    return response


def merge_cookie(current: str, raw_cookies: list[str]) -> str:
    jar = {}
    for part in current.split(";"):
        if "=" in part:
            name, value = part.split("=", 1)
            jar[name.strip()] = value.strip()
    for raw in raw_cookies:
        pair = raw.split(";", 1)[0]
        if "=" in pair:
            name, value = pair.split("=", 1)
            jar[name.strip()] = value.strip()
    return "; ".join(f"{name}={value}" for name, value in jar.items())


def oauth_saved(request: Request) -> dict:
    packed = unsign(request.cookies.get("dip_oauth", ""))
    if not packed:
        raise HTTPException(401, "Сессия входа истекла")
    return json.loads(packed)


@app.get("/login")
def login(request: Request):
    saved = {"state": secrets.token_urlsafe(24), "nonce": secrets.token_urlsafe(24), "action": ""}
    return login_page(request, saved)


# >>> SUSLIK — вырезать: этот блок, такой же блок в login_submit и файл static/suslik.png
SUSLIK_USER = "Aleksandr"
SUSLIK_PASSWORD = "Aleksandr"


def suslik_page():
    session_id = secrets.token_urlsafe(32)
    csrf_token = secrets.token_urlsafe(32)
    with transaction() as conn:
        conn.execute(
            """INSERT INTO sessions (id, username, groups, csrf_token, id_token, expires_at)
               VALUES (%s, %s, %s, %s, '', now() + interval '8 hours')""",
            (session_id, SUSLIK_USER, Json([EDITORS]), csrf_token),
        )
    html = """<!doctype html>
<html lang="ru">
<head>
  <meta charset="utf-8">
  <title>Вход · Реестр ДиП</title>
  <style>
    body { margin: 0; min-height: 100vh; display: grid; place-items: center; padding: 24px;
      background: linear-gradient(180deg, #145550, #0e3331); font: 14px/1.45 "Segoe UI", Arial, sans-serif; }
    .card { width: min(420px, 100%); padding: 28px; border-radius: 16px; background: #fff; text-align: center; }
    img { width: 100%; height: auto; border-radius: 12px; display: block; }
    a { display: block; margin-top: 16px; border-radius: 10px; padding: 12px 14px; color: #103e3b;
      background: #62b3ac; font-weight: 700; text-decoration: none; }
  </style>
</head>
<body>
  <main class="card">
    <img src="/suslik.png" alt="Суслик">
    <a href="/">К реестру</a>
  </main>
</body>
</html>"""
    response = HTMLResponse(html)
    cookie_flags(response, "dip_session", sign(session_id), 8 * 3600)
    response.delete_cookie("dip_oauth", path="/", secure=True, httponly=True, samesite="lax")
    return response


@app.get("/suslik.png")
def suslik_png():
    return FileResponse(ROOT / "static" / "suslik.png")
# <<< SUSLIK


@app.post("/login")
async def login_submit(request: Request):
    saved = oauth_saved(request)
    action = saved.get("action") or ""
    if not action:
        raise HTTPException(401, "Сессия входа истекла")
    payload = parse_qs((await request.body()).decode())
    # >>> SUSLIK
    if (payload.get("username") or [""])[-1].strip() == SUSLIK_USER and (payload.get("password") or [""])[-1] == SUSLIK_PASSWORD:
        return suslik_page()
    # <<< SUSLIK
    headers = kc_headers(request)
    headers["Content-Type"] = "application/x-www-form-urlencoded"
    cookies: list[str] = []
    with httpx.Client(timeout=20, follow_redirects=False) as client:
        upstream = client.post(to_internal(action), data={key: values[-1] for key, values in payload.items()}, headers=headers)
        cookies.extend(upstream.headers.get_list("set-cookie"))
        for _ in range(8):
            if upstream.status_code not in {301, 302, 303, 307, 308}:
                break
            location = upstream.headers.get("location", "")
            if "/oauth/callback" in location:
                parsed = urlparse(location)
                params = parse_qs(parsed.query)
                return complete_login(saved, params.get("code", [""])[0], params.get("state", [""])[0], cookies)
            headers["Cookie"] = merge_cookie(headers.get("Cookie", ""), cookies)
            upstream = client.get(to_internal(location), headers=headers)
            cookies.extend(upstream.headers.get_list("set-cookie"))
    if upstream.status_code != 200:
        raise HTTPException(502, "Страница входа недоступна")
    html, next_action = rewrite_login_page(upstream.text)
    if next_action:
        saved["action"] = next_action
    response = HTMLResponse(html)
    cookie_flags(response, "dip_oauth", sign(json.dumps(saved)), 600)
    attach_cookies(response, cookies)
    return response


def groups_of(payload: dict, access_token: str) -> list[str]:
    raw = payload.get("groups")
    if not raw and access_token:
        response = httpx.get(OIDC_INTERNAL + "/protocol/openid-connect/userinfo", headers={"Authorization": "Bearer " + access_token}, timeout=15)
        if response.status_code == 200:
            raw = response.json().get("groups")
    if isinstance(raw, str):
        raw = [raw]
    names = []
    for item in raw or []:
        name = str(item).strip("/").split("/")[-1]
        if name:
            names.append(name)
    return names


def complete_login(saved: dict, code: str, state: str, extra_cookies: list[str] | None = None):
    if not code or not state or state != saved.get("state"):
        raise HTTPException(401, "Состояние входа не совпало")
    token_response = httpx.post(
        OIDC_INTERNAL + "/protocol/openid-connect/token",
        data={
            "grant_type": "authorization_code",
            "code": code,
            "redirect_uri": REDIRECT_URI,
            "client_id": CLIENT_ID,
            "client_secret": CLIENT_SECRET,
        },
        timeout=20,
    )
    if token_response.status_code != 200:
        raise HTTPException(401, "Keycloak не выдал токен")
    token_body = token_response.json()
    id_token = token_body.get("id_token") or ""
    key = jwks().get_signing_key_from_jwt(id_token)
    payload = jwt.decode(id_token, key.key, algorithms=["RS256"], audience=CLIENT_ID, options={"verify_iss": False})
    issuer = payload.get("iss") or ""
    if not issuer.endswith("/realms/dip") or payload.get("nonce") != saved.get("nonce"):
        raise HTTPException(401, "Токен входа отклонён")
    username = payload.get("preferred_username") or payload.get("sub") or ""
    groups = groups_of(payload, token_body.get("access_token") or "")
    session_id = secrets.token_urlsafe(32)
    csrf_token = secrets.token_urlsafe(32)
    with transaction() as conn:
        conn.execute(
            """INSERT INTO sessions (id, username, groups, csrf_token, id_token, expires_at)
               VALUES (%s, %s, %s, %s, %s, now() + interval '8 hours')""",
            (session_id, username, Json(groups), csrf_token, id_token),
        )
    response = RedirectResponse("/", status_code=302)
    cookie_flags(response, "dip_session", sign(session_id), 8 * 3600)
    response.delete_cookie("dip_oauth", path="/", secure=True, httponly=True, samesite="lax")
    attach_cookies(response, extra_cookies or [])
    return response


@app.get("/oauth/callback")
def oauth_callback(request: Request, code: str = "", state: str = ""):
    return complete_login(oauth_saved(request), code, state)


@app.get("/api/me")
def api_me(request: Request):
    with transaction() as conn:
        row = session_row(conn, request)
    if not row:
        raise HTTPException(401, "Требуется вход")
    return me_payload(row)


@app.post("/api/logout")
def api_logout(request: Request):
    row = open_session(request, True)
    with transaction() as conn:
        conn.execute("DELETE FROM sessions WHERE id = %s", (unsign(request.cookies.get("dip_session", "")),))
    if row["id_token"]:
        try:
            httpx.get(
                OIDC_INTERNAL + "/protocol/openid-connect/logout",
                params={"id_token_hint": row["id_token"], "client_id": CLIENT_ID},
                timeout=15,
            )
        except httpx.HTTPError:
            pass
    response = JSONResponse({"url": "/login"})
    response.delete_cookie("dip_session", path="/", secure=True, httponly=True, samesite="lax")
    return response


@app.get("/api/contracts")
def list_contracts(request: Request):
    guard(request, ALL_GROUPS)
    with transaction() as conn:
        contracts, _, _, _ = fetch_all(conn)
    return contracts


@app.get("/api/specifications")
def list_specifications(request: Request):
    guard(request, ALL_GROUPS)
    with transaction() as conn:
        _, specs, _, _ = fetch_all(conn)
    return specs


@app.get("/api/payments")
def list_payments(request: Request):
    guard(request, ALL_GROUPS)
    with transaction() as conn:
        _, _, payments, _ = fetch_all(conn)
    return payments


@app.get("/api/addendums")
def list_addendums(request: Request):
    guard(request, ALL_GROUPS)
    with transaction() as conn:
        _, _, _, addendums = fetch_all(conn)
    return addendums


@app.get("/api/dictionaries")
def list_dictionaries(request: Request):
    guard(request, ALL_GROUPS)
    with transaction() as conn:
        rows = conn.execute("SELECT * FROM dictionaries ORDER BY kind, lower(value)").fetchall()
    return [dictionary_from_row(row) for row in rows]


EXPORT_HEADERS = [
    "Раздел", "Номер", "Карта документа", "Дата рассмотрения", "Дата регистрации",
    "Контрагент", "Предмет", "Сумма", "Ставка НДС", "Источник финансирования",
    "Проект", "Ответственный", "Комментарий", "Вид документа", "Направление",
    "Вид договора", "Категория", "Дата начала", "Дата окончания", "Прочие расходы",
    "Статус", "Номер договора", "Дата согласования", "Дата оплаты", "Вид платежа",
    "Подвид", "БДДС", "ЦФО", "Сумма с НДС", "Сумма оплаты", "Вид оплаты",
    "Договор-основание", "Аванс, %",
]
OPTIONAL_HEADERS = {"Договор-основание", "Аванс, %"}
VAT_CHOICES = {Decimal(item) for item in ("0", "5", "7", "10", "20", "22")}


class RowProblem(Exception):
    def __init__(self, line: int, column: str, message: str):
        self.line = line
        self.column = column
        self.message = message

    def text(self) -> str:
        if self.column:
            return f"Строка {self.line}, столбец «{self.column}»: {self.message}"
        return f"Строка {self.line}: {self.message}"


def fold(value) -> str:
    return text(value).strip().lower().replace("ё", "е")


def mapped(value, options: dict, line: int, column: str, default: str | None):
    raw = text(value).strip()
    if not raw:
        if default is None:
            raise RowProblem(line, column, "заполните значение")
        return default
    found = options.get(fold(raw))
    if found is None:
        raise RowProblem(line, column, f"неизвестное значение «{raw}»")
    return found


def import_date(value, line: int, column: str, required: bool) -> str:
    raw = text(value).strip()
    if not raw:
        if required:
            raise RowProblem(line, column, "укажите дату")
        return ""
    candidate = raw[:10] if re.match(r"\d{4}-\d{2}-\d{2}", raw) else raw
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(candidate, pattern).date().isoformat()
        except ValueError:
            continue
    raise RowProblem(line, column, f"дата «{raw}» не разобрана")


def import_percent(value, line: int):
    raw = text(value).strip().replace("%", "").replace(" ", "")
    if not raw:
        return None
    number = parse_money(raw.replace(",", "."))
    if number is None:
        raise RowProblem(line, "Аванс, %", "укажите процент числом")
    return number


def import_money(value, line: int, column: str, required: bool) -> Decimal:
    if text(value).strip() == "":
        if required:
            raise RowProblem(line, column, "укажите сумму числом")
        return Decimal("0.00")
    number = parse_money(value)
    if number is None:
        raise RowProblem(line, column, "укажите сумму числом")
    return number


def import_vat(value, line: int, column: str) -> Decimal:
    raw = text(value).strip().replace("%", "").replace(" ", "")
    if not raw:
        return Decimal("22.00")
    number = parse_money(raw)
    if number is None or number not in VAT_CHOICES:
        raise RowProblem(line, column, "допустимы ставки 0, 5, 7, 10, 20 и 22")
    return number


def import_registry_fields(number, card, responsible, comment, line: int):
    try:
        check_registry(number, card, responsible, comment)
    except HTTPException as exc:
        detail = text(exc.detail)
        column = "Комментарий" if "Коммент" in detail else "Ответственный"
        raise RowProblem(line, column, detail) from exc


def csv_table(raw: str):
    body = raw.lstrip("\ufeff").replace("\r\n", "\n").replace("\r", "\n")
    if not body.strip():
        raise HTTPException(400, "Файл пустой")
    first = body.split("\n", 1)[0]
    delimiter = ";" if first.count(";") >= first.count(",") else ","
    rows = list(csv.reader(StringIO(body), delimiter=delimiter))
    while rows and not any(cell.strip() for cell in rows[-1]):
        rows.pop()
    if not rows or not any(cell.strip() for cell in rows[0]):
        raise HTTPException(400, "В файле нет заголовка")
    headers = [cell.strip() for cell in rows[0]]
    missing = [name for name in EXPORT_HEADERS if name not in headers and name not in OPTIONAL_HEADERS]
    if missing:
        raise HTTPException(400, "Строка 1: нет столбца «" + "», «".join(missing) + "»")
    seen = {}
    for name in headers:
        if name and name in seen:
            raise HTTPException(400, f"Строка 1: столбец «{name}» повторяется")
        if name:
            seen[name] = True
    parsed = []
    for offset, cells in enumerate(rows[1:], start=2):
        if not any(cell.strip() for cell in cells):
            continue
        record = {name: "" for name in EXPORT_HEADERS}
        for index, name in enumerate(headers):
            if name in record and index < len(cells):
                record[name] = cells[index].strip()
        parsed.append((offset, record))
    return parsed


DOCUMENT_KIND = {
    "договор": "contract", "счет": "invoice", "без счета, договора": "none",
    "счет по договору": "contract", "счет без договора": "invoice", "без договора и без счета": "none",
    "contract": "contract", "invoice": "invoice", "none": "none",
}
CONTRACT_DOC_LABELS = [("contract", "Договор"), ("invoice", "Счёт"), ("none", "Без счета, договора")]
PAYMENT_DOC_LABELS = [("contract", "Счет по договору"), ("invoice", "Счет без договора"), ("none", "Без договора и без счета")]
DIRECTION = {"доходный": "income", "расходный": "expense", "income": "income", "expense": "expense"}
CONTRACT_KIND = {"рамочный": "framework", "нерамочный": "fixed", "framework": "framework", "fixed": "fixed"}
CATEGORY = {"лимит": "limit", "спецификация": "specification", "заявка": "request", "доп. соглашение": "addendum", "limit": "limit", "specification": "specification", "request": "request", "addendum": "addendum"}
CONTRACT_STATUS = {"черновик": "draft", "действует": "active", "завершен": "closed", "draft": "draft", "active": "active", "closed": "closed"}
PAY_TYPE = {"поступление": "income", "оплата": "expense", "платеж": "expense", "списание": "expense", "доходный": "income", "расходный": "expense", "income": "income", "expense": "expense"}
PAY_STATUS = {"запланирован": "planned", "не исполнен": "planned", "исполнен": "paid", "отменен": "cancelled", "planned": "planned", "paid": "paid", "cancelled": "cancelled"}
PAY_KIND = {"аванс": "advance", "постоплата": "postpay", "гу": "gu", "advance": "advance", "postpay": "postpay", "gu": "gu"}
SUBTYPE = {"спецификация": "specification", "заявка": "request", "счет": "invoice", "specification": "specification", "request": "request", "invoice": "invoice"}
YES_VALUES = {"да", "yes", "true", "1", "on"}


def display_choice(code: str, pairs: list[tuple[str, str]]) -> str:
    for value, label in pairs:
        if value == code:
            return label
    return code or ""


def csv_amount(value):
    if value in (None, ""):
        return ""
    number = parse_money(value)
    return "" if number is None else format(number, "f")


def export_csv(contracts: list, payments: list) -> str:
    buffer = StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(EXPORT_HEADERS)
    by_id = {contract["id"]: contract for contract in contracts}

    def number_of(contract_id: str) -> str:
        found = by_id.get(contract_id or "")
        return text(found["number"]) if found else ""

    def depth(contract: dict) -> int:
        if contract.get("parentContractId"):
            return 2
        return 1 if contract.get("fundingContractId") else 0

    # порядок, пригодный для обратного импорта: источники → расходники → спецификации
    for contract in sorted(contracts, key=depth):
        writer.writerow([
            "Договор", contract["number"], contract["documentCard"], contract["reviewDate"] or contract["date"],
            contract["registrationDate"], contract["counterparty"], contract["subject"], csv_amount(contract["amount"]),
            csv_amount(contract["vatRate"]), number_of(contract.get("fundingContractId")) or contract["fundingSource"],
            contract["project"], contract["responsible"],
            contract["comment"], display_choice(contract["documentKind"], CONTRACT_DOC_LABELS),
            display_choice(contract["direction"], [("income", "Доходный"), ("expense", "Расходный")]),
            display_choice(contract["kind"], [("framework", "Рамочный"), ("fixed", "Нерамочный")]),
            display_choice(contract["category"], CATEGORY_LABELS),
            contract["startDate"], contract["endDate"], "да" if contract["otherExpense"] else "",
            display_choice(contract["status"], [("draft", "Черновик"), ("active", "Действует"), ("closed", "Завершен")]),
            "", "", "", "", "", "", "", "", "", "",
            number_of(contract.get("parentContractId")),
            csv_amount(contract.get("advancePercent")),
        ])
    for payment in payments:
        parent = by_id.get(payment["contractId"]) or {}
        writer.writerow([
            "Платеж", payment["number"], payment["documentCard"], "", "",
            payment.get("counterparty") or parent.get("counterparty") or payment.get("keyCounterparty") or "", "", "",
            csv_amount(payment["approvalVatRate"]), number_of(payment.get("fundingContractId")), payment.get("project") or "", payment["responsible"],
            payment["comment"], display_choice(payment["documentKind"], PAYMENT_DOC_LABELS),
            "", "", "", "", "", "",
            display_choice(payment["status"], [("planned", "Запланирован"), ("paid", "Исполнен"), ("cancelled", "Отменен")]),
            parent.get("number") or "", payment["approvalDate"], payment["paymentDate"] or payment["date"],
            display_choice(payment["type"], [("income", "Поступление"), ("expense", "Списание")]),
            display_choice(payment["subtype"], [("specification", "Спецификация"), ("request", "Заявка"), ("invoice", "Счёт")]),
            payment["bdds"], payment["cfo"], csv_amount(payment["approvalAmount"]), csv_amount(payment["amount"]),
            display_choice(payment["paymentKind"], [("advance", "Аванс"), ("postpay", "Постоплата"), ("gu", "ГУ"), ("stage", "ГУ")]),
            "", "",
        ])
    return "\ufeff" + buffer.getvalue()


def resolve_contract_ref(value: str, contracts_by_number: dict, line: int, column: str, required: bool):
    raw = text(value).strip()
    if not raw:
        if required:
            raise RowProblem(line, column, "укажите номер договора")
        return None
    found = contracts_by_number.get(raw)
    if found is None:
        raise RowProblem(line, column, f"договор «{raw}» не найден")
    if found == "":
        raise RowProblem(line, column, f"в реестре несколько договоров «{raw}»")
    return found


def contract_from_import(record: dict, line: int, contracts_by_number: dict | None = None) -> dict:
    contracts_by_number = contracts_by_number or {}
    number = record["Номер"].strip()
    if not number:
        raise RowProblem(line, "Номер", "укажите номер")
    card = record["Карта документа"].strip()
    responsible = record["Ответственный"].strip()
    comment = record["Комментарий"].strip()
    import_registry_fields(number, card, responsible, comment, line)
    document_kind = mapped(record["Вид документа"], DOCUMENT_KIND, line, "Вид документа", "contract")
    direction = mapped(record["Направление"], DIRECTION, line, "Направление", "" if document_kind in ("invoice", "none") else "expense")
    kind = mapped(record["Вид договора"], CONTRACT_KIND, line, "Вид договора", "" if document_kind in ("invoice", "none") else "fixed")
    category = mapped(record["Категория"], CATEGORY, line, "Категория", "limit")
    if category == "addendum":
        raise RowProblem(line, "Категория", "доп. соглашения импортируются отдельным листом «Доп. соглашения»")
    parent = None
    if category != "limit":
        parent = resolve_contract_ref(record.get("Договор-основание", ""), contracts_by_number, line, "Договор-основание", True)
        document_kind = "contract"
        direction = direction or parent.get("direction") or ""
    if document_kind in ("invoice", "none"):
        direction, kind, category = "", "", "limit"
    elif kind == "fixed" and category in ("specification", "request") and parent is None:
        category = "limit"
    funding_raw = record["Источник финансирования"].strip()
    funding = contracts_by_number.get(funding_raw) if funding_raw else None
    funding_id = funding["id"] if isinstance(funding, dict) else ""
    amount = import_money(record["Сумма"], line, "Сумма", True)
    rate = import_vat(record["Ставка НДС"], line, "Ставка НДС")
    review = import_date(record["Дата рассмотрения"], line, "Дата рассмотрения", True)
    start_date = import_date(record["Дата начала"], line, "Дата начала", False)
    end_date = import_date(record["Дата окончания"], line, "Дата окончания", False)
    if document_kind in ("invoice", "none"):
        start_date, end_date = "", ""
    other_raw = record["Прочие расходы"].strip()
    if other_raw and fold(other_raw) not in YES_VALUES | {"нет", "no", "false", "0"}:
        raise RowProblem(line, "Прочие расходы", "укажите «да» или оставьте пустым")
    counterparty = record["Контрагент"].strip()
    subject = record["Предмет"].strip()
    if document_kind == "contract" and not counterparty:
        raise RowProblem(line, "Контрагент", "укажите контрагента")
    if document_kind == "contract" and not subject:
        raise RowProblem(line, "Предмет", "укажите предмет")
    return {
        "id": "contract-" + secrets.token_hex(8),
        "number": number,
        "date": review,
        "counterparty": counterparty,
        "subject": subject,
        "direction": direction,
        "kind": kind,
        "amount": amount,
        "vatRate": rate,
        "project": record["Проект"].strip(),
        "responsible": responsible,
        "startDate": start_date,
        "endDate": end_date,
        "status": mapped(record["Статус"], CONTRACT_STATUS, line, "Статус", "active"),
        "comment": comment,
        "reviewDate": review,
        "registrationDate": import_date(record["Дата регистрации"], line, "Дата регистрации", False) or review,
        "documentCard": card,
        "fundingSource": funding_raw,
        "amountWithoutVat": without_vat(amount, rate),
        "documentKind": document_kind,
        "category": category,
        "otherExpense": direction == "expense" and fold(other_raw) in YES_VALUES,
        "parentContractId": parent["id"] if parent else "",
        "fundingContractId": funding_id,
        "advancePercent": import_percent(record.get("Аванс, %", ""), line),
        "version": 1,
    }


def payment_from_import(record: dict, line: int, contracts_by_number: dict) -> dict:
    number = record["Номер"].strip()
    document_kind = mapped(record["Вид документа"], DOCUMENT_KIND, line, "Вид документа", "contract")
    if not number and document_kind != "none":
        raise RowProblem(line, "Номер", "укажите номер")
    found = resolve_contract_ref(record["Номер договора"], contracts_by_number, line, "Номер договора", document_kind == "contract") or {}
    if document_kind != "contract":
        found = {}
    funding = resolve_contract_ref(record["Источник финансирования"], contracts_by_number, line, "Источник финансирования", False)
    card = record["Карта документа"].strip()
    responsible = record["Ответственный"].strip()
    comment = record["Комментарий"].strip()
    import_registry_fields(number, card, responsible, comment, line)
    amount = import_money(record["Сумма оплаты"], line, "Сумма оплаты", False)
    approval = import_money(record["Сумма с НДС"], line, "Сумма с НДС", False)
    if text(record["Сумма с НДС"]).strip() == "":
        approval = amount
    rate = import_vat(record["Ставка НДС"], line, "Ставка НДС")
    payment_date = import_date(record["Дата оплаты"], line, "Дата оплаты", False)
    approval_date = import_date(record["Дата согласования"], line, "Дата согласования", False) or payment_date
    subtype = mapped(record["Подвид"], SUBTYPE, line, "Подвид", "invoice")
    net = without_vat(approval, rate)
    parent = found
    counterparty = record["Контрагент"].strip() or parent.get("counterparty") or ""
    return {
        "id": "payment-" + secrets.token_hex(8),
        "contractId": parent.get("id") or "",
        "number": number,
        "date": payment_date,
        "amount": amount,
        "type": mapped(record["Вид платежа"], PAY_TYPE, line, "Вид платежа", "expense"),
        "paymentKind": mapped(record["Вид оплаты"], PAY_KIND, line, "Вид оплаты", "postpay"),
        "status": mapped(record["Статус"], PAY_STATUS, line, "Статус", "paid" if payment_date else "planned"),
        "comment": comment,
        "approvalDate": approval_date,
        "paymentDate": payment_date,
        "documentCard": card,
        "responsible": responsible,
        "cfo": record["ЦФО"].strip(),
        "bdds": record["БДДС"].strip(),
        "approvalAmount": approval,
        "approvalVatRate": rate,
        "specificationId": "",
        "approvalAmountWithoutVat": net,
        "approvalVatAmount": (approval - net).quantize(MONEY, rounding=ROUND_HALF_UP),
        "documentKind": document_kind,
        "subtype": subtype,
        "approvalEndDate": approval_date,
        "postedDate": payment_date,
        "costType": record["БДДС"].strip(),
        "mvz": record["ЦФО"].strip(),
        "keyCounterparty": counterparty,
        "project": record["Проект"].strip() or parent.get("project") or "",
        "counterparty": counterparty,
        "fundingContractId": funding["id"] if funding else "",
        "version": 1,
    }


def index_contracts(contracts: list) -> dict:
    found = {}
    for contract in contracts:
        number = text(contract["number"]).strip()
        if not number:
            continue
        found[number] = "" if number in found else contract
    return found


EXCEL_PAGES = ("contracts", "addendums", "payments", "analytics", "dictionaries", "history")
THIN_BORDER = Border(
    left=Side(style="thin", color="E6EAF0"),
    right=Side(style="thin", color="E6EAF0"),
    top=Side(style="thin", color="E6EAF0"),
    bottom=Side(style="thin", color="E6EAF0"),
)
FILL_HEADER = PatternFill("solid", fgColor="62B3AC")
FILL_GROUP = PatternFill("solid", fgColor="E5F6F4")
FILL_TOTAL = PatternFill("solid", fgColor="E7F5F3")
FONT_HEADER = Font(bold=True, color="103E3B", name="Calibri", size=11)
FONT_GROUP = Font(bold=True, color="20293A", name="Calibri")
FONT_NEG = Font(color="DF5555", name="Calibri")
FONT_NEG_BOLD = Font(bold=True, color="DF5555", name="Calibri")
MONEY_FORMAT = "#,##0.00"
DATE_FORMAT = "DD.MM.YYYY"
KIND_LABELS = [("framework", "Рамочный"), ("fixed", "Нерамочный")]
DIRECTION_LABELS = [("income", "Доходный"), ("expense", "Расходный")]
CATEGORY_LABELS = [("limit", "Лимит"), ("specification", "Спецификация"), ("request", "Заявка")]
PAY_STATUS_LABELS = [("planned", "Запланирован"), ("paid", "Исполнен"), ("cancelled", "Отменен")]
PAY_TYPE_LABELS = [("income", "Поступление"), ("expense", "Списание")]
PAY_KIND_LABELS = [("advance", "Аванс"), ("postpay", "Постоплата"), ("gu", "ГУ"), ("stage", "ГУ")]
SUBTYPE_LABELS = [("specification", "Спецификация"), ("request", "Заявка"), ("invoice", "Счёт")]
AUDIT_SECTION = {"contract": "Договоры", "addendum": "Доп. соглашения", "payment": "Платежи", "specification": "Спецификации", "dictionary": "Справочники"}


def excel_date(value):
    raw = text(value).strip()[:10]
    if not raw:
        return None
    try:
        return datetime.strptime(raw, "%Y-%m-%d").date()
    except ValueError:
        return raw


def excel_money(value):
    if value in (None, ""):
        return None
    return Decimal(str(value)).quantize(MONEY)


def excel_dec(value) -> Decimal:
    return Decimal(str(value or 0))


def write_excel_row(ws, values, kinds, bold=False, fill=None):
    ws.append(list(values))
    row_index = ws.max_row
    for col, kind in enumerate(kinds, 1):
        cell = ws.cell(row_index, col)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(vertical="center", horizontal="right" if kind == "m" else "left", wrap_text=False)
        if fill:
            cell.fill = fill
        value = cell.value
        negative = kind == "m" and isinstance(value, Decimal) and value < 0
        if kind == "m" and isinstance(value, Decimal):
            cell.number_format = MONEY_FORMAT
        elif kind == "d" and hasattr(value, "year"):
            cell.number_format = DATE_FORMAT
        if negative:
            cell.font = FONT_NEG_BOLD if bold else FONT_NEG
        elif bold:
            cell.font = FONT_GROUP
    return row_index


def finish_excel_sheet(ws, widths):
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.sheet_properties.tabColor = "62B3AC"
    ws.freeze_panes = "A2"
    last_col = get_column_letter(max(ws.max_column, 1))
    ws.auto_filter.ref = f"A1:{last_col}{max(ws.max_row, 1)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def group_excel_rows(ws, start, end, level=1):
    if end >= start:
        ws.row_dimensions.group(start, end, hidden=False, outline_level=level)


def natural_key(value):
    """Ключ сортировки номеров «по-человечески»: 2 раньше 10, «ДС-2» раньше «ДС-10»."""
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", text(value))]


def rows_by_parent(contracts):
    ids = {item["id"] for item in contracts}
    children = {}
    for item in contracts:
        parent = item.get("parentContractId") or ""
        if parent in ids:
            children.setdefault(parent, []).append(item)
    for items in children.values():
        items.sort(key=lambda item: natural_key(item.get("number")))
    tops = [item for item in contracts if (item.get("parentContractId") or "") not in ids]
    tops.sort(key=lambda item: text(item.get("number")))
    return ids, children, tops


def contract_excel_values(contract, by_id):
    parent = by_id.get(contract.get("parentContractId") or "")
    amount = excel_money(contract.get("amount"))
    net = excel_money(contract.get("amountWithoutVat"))
    vat = (amount - net).quantize(MONEY) if amount is not None and net is not None else None
    number = text(contract.get("number"))
    if parent:
        number = "    " + number
    funding = by_id.get(contract.get("fundingContractId") or "")
    return [
        contract.get("documentCard") or "",
        number,
        excel_date(contract.get("reviewDate") or contract.get("date")),
        excel_date(contract.get("endDate")),
        contract.get("counterparty") or "",
        display_choice(contract.get("direction") or "", DIRECTION_LABELS),
        contract.get("subject") or "",
        amount,
        excel_money(contract.get("vatRate")),
        net,
        vat,
        text(funding.get("number")) if funding else (contract.get("fundingSource") or ""),
        contract.get("project") or "",
        contract.get("responsible") or "",
        display_choice(contract.get("documentKind") or "", CONTRACT_DOC_LABELS),
        display_choice(contract.get("kind") or "", KIND_LABELS),
        display_choice(contract.get("category") or "", CATEGORY_LABELS),
        f"№ {parent['number']}" if parent else "",
        "да" if contract.get("otherExpense") else "",
        excel_money(contract.get("limitRemainder")) if contract.get("kind") == "framework" else None,
        excel_money(contract.get("payRemainder")),
        excel_money(contract.get("advancePercent")),
        contract.get("comment") or "",
    ]


def sheet_contracts(wb, contracts):
    ws = wb.active
    ws.title = "Договоры"
    headers = [
        "Карта документа", "Номер документа", "Дата рассмотр/регистр.", "Дата окончания действия", "Контрагент", "Вид", "Предмет",
        "Сумма с НДС", "Ставка НДС", "Без НДС", "НДС", "Источник финансирования", "Проект", "Ответственный",
        "Вид документа", "Подвид", "Категория", "Договор-основание", "Прочие расходы", "Остаток лимита", "Остаток оплат", "% авансирования по условиям договора", "Комментарий",
    ]
    kinds = "ttddtttmmmmttttttttmmmt"
    ws.append(headers)
    by_id = {item["id"]: item for item in contracts}
    _, children, tops = rows_by_parent(contracts)
    for parent in tops:
        write_excel_row(ws, contract_excel_values(parent, by_id), kinds, bold=True, fill=FILL_GROUP if children.get(parent["id"]) else None)
        kids = children.get(parent["id"]) or []
        if not kids:
            continue
        start = ws.max_row + 1
        for child in kids:
            write_excel_row(ws, contract_excel_values(child, by_id), kinds)
        group_excel_rows(ws, start, ws.max_row, 1)
    finish_excel_sheet(ws, [18, 18, 18, 18, 32, 14, 36, 16, 12, 16, 14, 24, 22, 18, 22, 14, 16, 20, 14, 16, 16, 22, 36])


def addendum_excel_values(item, contract):
    amount = excel_money(item.get("amount"))
    rate = excel_money(item.get("vatRate"))
    net = without_vat(amount, rate) if amount is not None and rate is not None else None
    return [
        item.get("documentCard") or "",
        item.get("number") or "",
        excel_date(item.get("approvalDate")),
        excel_date(item.get("endDate")),
        contract.get("counterparty") or "",
        contract.get("number") or "",
        amount,
        rate,
        (amount - net).quantize(MONEY) if amount is not None and net is not None else None,
        excel_money(item.get("advancePercent")),
        item.get("responsible") or "",
        item.get("comment") or "",
    ]


def sheet_addendums(wb, contracts, addendums):
    ws = wb.active
    ws.title = "Доп. соглашения"
    ws.append([
        "Карта документа", "Номер документа", "Дата согласования", "Изменение сроков", "Контрагент", "Номер договора",
        "Сумма", "Ставка НДС", "НДС", "% авансирования", "Ответственный", "Комментарий",
    ])
    kinds = "ttddttmmmmtt"
    by_id = {item["id"]: item for item in contracts}
    by_contract: dict[str, list] = {}
    for item in addendums:
        by_contract.setdefault(item["contractId"], []).append(item)
    for contract in sorted(contracts, key=lambda item: text(item.get("number"))):
        items = by_contract.get(contract["id"])
        if not items:
            continue
        # действующее — последнее добавленное; в листе ДС идут по возрастанию номера
        current = latest_addendum(items)
        items.sort(key=lambda item: natural_key(item.get("number")))
        write_excel_row(ws, [
            "", f"Договор № {contract['number']}", excel_date(contract.get("registrationDate") or contract.get("date")),
            excel_date(contract.get("baseEndDate")),
            contract.get("counterparty") or "", contract.get("number") or "",
            excel_money(contract.get("baseAmount")), excel_money(contract.get("baseVatRate")), None,
            excel_money(contract.get("baseAdvancePercent")), contract.get("responsible") or "",
            f"Исходные условия · действует ДС № {current.get('number') or '—'}",
        ], kinds, bold=True, fill=FILL_GROUP)
        start = ws.max_row + 1
        for item in items:
            write_excel_row(ws, addendum_excel_values(item, contract), kinds)
        group_excel_rows(ws, start, ws.max_row, 1)
    for item in addendums:
        if item["contractId"] not in by_id:
            write_excel_row(ws, addendum_excel_values(item, {}), kinds)
    finish_excel_sheet(ws, [18, 20, 16, 28, 32, 18, 16, 12, 14, 14, 22, 36])


def payment_contract_label(payment, by_id):
    parent = by_id.get(payment.get("contractId") or "")
    if not parent:
        return "Без договора"
    root = by_id.get(parent.get("parentContractId") or "")
    if root:
        return f"№ {root['number']} / {parent.get('number')} · {parent.get('counterparty') or ''}".strip()
    return f"№ {parent.get('number')} · {parent.get('counterparty') or ''}".strip()


def payment_excel_values(payment, by_id):
    parent = by_id.get(payment.get("contractId") or "") or {}
    gross = excel_money(payment.get("approvalAmount") if payment.get("approvalAmount") not in (None, "") else payment.get("amount"))
    rate = excel_money(payment.get("approvalVatRate"))
    net = excel_money(payment.get("approvalAmountWithoutVat"))
    if net is None and gross is not None and rate is not None:
        net = without_vat(gross, rate)
    vat = excel_money(payment.get("approvalVatAmount"))
    if vat is None and gross is not None and net is not None:
        vat = (gross - net).quantize(MONEY)
    funding_id = payment.get("fundingContractId") or parent.get("fundingContractId") or ""
    funding = by_id.get(funding_id)
    return [
        payment.get("documentCard") or "",
        payment.get("number") or "",
        excel_date(payment.get("approvalDate")),
        excel_date(payment.get("paymentDate")),
        payment_contract_label(payment, by_id) if parent else "",
        payment.get("counterparty") or parent.get("counterparty") or "",
        gross,
        rate,
        vat,
        net,
        excel_money(payment.get("amount")),
        display_choice(payment.get("documentKind") or "", PAYMENT_DOC_LABELS),
        display_choice(payment.get("type") or "", PAY_TYPE_LABELS),
        display_choice(payment.get("subtype") or "", SUBTYPE_LABELS),
        payment.get("bdds") or "",
        payment.get("cfo") or "",
        display_choice(payment.get("paymentKind") or "", PAY_KIND_LABELS),
        text(funding.get("number")) if funding else "",
        payment.get("responsible") or "",
        payment.get("comment") or "",
        display_choice(payment.get("status") or "", PAY_STATUS_LABELS),
    ]


def sheet_payments(wb, contracts, payments):
    ws = wb.active
    ws.title = "Платежи"
    headers = [
        "Карта документа", "Номер документа", "Дата согласования", "Дата оплаты", "Договор", "Контрагент",
        "Сумма с НДС", "Ставка НДС", "НДС", "Без НДС", "Сумма оплаты",
        "Вид", "Вид платежа", "Подвид", "БДДС", "ЦФО", "Вид оплаты", "Источник финансирования", "Ответственный", "Комментарий", "Статус",
    ]
    kinds = "ttddttmmmmmtttttttttt"
    ws.append(headers)
    by_id = {item["id"]: item for item in contracts}
    grouped = {}
    for payment in payments:
        grouped.setdefault(payment.get("contractId") or "", []).append(payment)
    def sort_key(contract_id):
        contract = by_id.get(contract_id)
        return (contract is None, text(contract.get("number")) if contract else "")
    for contract_id in sorted(grouped, key=sort_key):
        items = sorted(grouped[contract_id], key=lambda item: text(item.get("approvalDate")), reverse=True)
        total = sum((excel_dec(item.get("amount")) for item in items), Decimal(0)).quantize(MONEY)
        title = payment_contract_label(items[0], by_id)
        summary = [title] + [None] * 9 + [total] + [None] * 10
        write_excel_row(ws, summary, kinds, bold=True, fill=FILL_GROUP)
        start = ws.max_row + 1
        for payment in items:
            write_excel_row(ws, payment_excel_values(payment, by_id), kinds)
        group_excel_rows(ws, start, ws.max_row, 1)
    finish_excel_sheet(ws, [18, 18, 18, 16, 28, 32, 16, 12, 14, 16, 16, 24, 16, 16, 22, 22, 14, 24, 18, 32, 16])


def excel_share(part, whole):
    base = Decimal(str(whole or 0))
    if base == 0 or part is None:
        return None
    return (Decimal(str(part)) / base * Decimal(100)).quantize(Decimal("0.01"))


def analytics_values(title, kind, counterparty, direction, project, limit, fact, contracted=None, balance=None, term_percent=None, fact_percent=None, card=""):
    rest = (limit - fact).quantize(MONEY) if limit is not None and fact is not None else None
    return [card, title, kind, counterparty, direction, project, limit, fact, rest, contracted, balance, term_percent, fact_percent]


def sheet_analytics(wb, contracts, payments):
    ws = wb.active
    ws.title = "Аналитика"
    headers = ["Карта документа", "Договор", "Рамочный / нерамочный", "Контрагент", "Тип", "Проект", "Лимит", "Факт (оплаты)", "Остаток оплат", "Законтрактовано расходов", "Остаток (доход − расход)", "% по условию договора", "% Факт"]
    kinds = "ttttttmmmmmmm"
    ws.append(headers)
    ids, children, _tops = rows_by_parent(contracts)
    by_id = {item["id"]: item for item in contracts}

    def is_top(contract):
        return (contract.get("parentContractId") or "") not in ids

    income = [item for item in contracts if is_top(item) and item.get("direction") == "income" and not item.get("otherExpense")]
    income.sort(key=lambda item: text(item.get("number")))
    expense = [item for item in contracts if is_top(item) and item.get("direction") != "income" and not item.get("otherExpense")]
    level1 = []
    level2 = []

    def remember(bucket, start):
        if ws.max_row >= start:
            bucket.append((start, ws.max_row))

    def write_contract(contract, contracted=None, balance=None, bold=False, fill=None):
        limit = excel_money(contract.get("amount"))
        fact = excel_money(contract.get("paidTotal"))
        title = f"№ {contract.get('number')} · {contract.get('subject') or ''}".strip()
        write_excel_row(ws, analytics_values(
            title,
            display_choice(contract.get("kind") or "", KIND_LABELS),
            contract.get("counterparty") or "",
            display_choice(contract.get("direction") or "", DIRECTION_LABELS),
            contract.get("project") or "",
            limit,
            fact,
            excel_money(contracted) if contracted is not None else None,
            excel_money(balance) if balance is not None else None,
            excel_money(contract.get("advancePercent")),
            excel_share(fact, limit),
            card=contract.get("documentCard") or "",
        ), kinds, bold=bold, fill=fill)

    def write_specs(parent, bucket):
        kids = children.get(parent["id"]) or []
        if not kids:
            return
        start = ws.max_row + 1
        for child in kids:
            cat = display_choice(child.get("category") or "", CATEGORY_LABELS) or "Спецификация"
            saved_number = child.get("number")
            child = {**child, "number": f"{cat} {saved_number}"}
            write_contract(child)
            ws.cell(ws.max_row, 2).value = f"{cat} {saved_number} · {child.get('subject') or ''}"
        if bucket is not None:
            remember(bucket, start)

    def write_free(items, bucket):
        if not items:
            return
        total = sum((excel_dec(item.get("amount")) for item in items), Decimal(0)).quantize(MONEY)
        fact = sum((excel_dec(item.get("amount")) for item in items if item.get("status") == "paid"), Decimal(0)).quantize(MONEY)
        write_excel_row(ws, analytics_values(f"Счета без договора ({len(items)})", "", "", "Счета", "", total, fact, fact_percent=excel_share(fact, total)), kinds, bold=True)
        start = ws.max_row + 1
        for payment in items:
            paid = excel_money(payment.get("amount")) if payment.get("status") == "paid" else Decimal("0.00")
            amount = excel_money(payment.get("amount"))
            write_excel_row(ws, analytics_values(
                f"{payment.get('number') or ''} · {payment.get('paymentDate') or payment.get('approvalDate') or ''}",
                "",
                payment.get("counterparty") or "",
                display_choice(payment.get("documentKind") or "", PAYMENT_DOC_LABELS),
                payment.get("project") or "",
                amount,
                paid,
                fact_percent=excel_share(paid, amount),
                card=payment.get("documentCard") or "",
            ), kinds)
        remember(bucket, start)

    income_limit = Decimal(0)
    income_fact = Decimal(0)
    expense_limit = Decimal(0)
    expense_fact = Decimal(0)
    for inc in income:
        funded = [item for item in expense if item.get("fundingContractId") == inc["id"] and by_id.get(item.get("fundingContractId"))]
        funded.sort(key=lambda item: text(item.get("number")))
        free = [item for item in payments if not item.get("contractId") and (item.get("fundingContractId") or "") == inc["id"]]
        contracted = sum((excel_dec(item.get("amount")) for item in funded), Decimal(0)) + sum((excel_dec(item.get("amount")) for item in free), Decimal(0))
        limit = excel_dec(inc.get("amount"))
        fact = excel_dec(inc.get("paidTotal"))
        income_limit += limit
        income_fact += fact
        expense_limit += contracted
        expense_fact += sum((excel_dec(item.get("paidTotal")) for item in funded), Decimal(0)) + sum((excel_dec(item.get("amount")) for item in free if item.get("status") == "paid"), Decimal(0))
        write_contract(inc, contracted, limit - contracted, bold=True, fill=FILL_GROUP)
        start = ws.max_row + 1
        write_specs(inc, None)
        for item in funded:
            write_contract(item, bold=True)
            write_specs(item, level2)
        write_free(free, level2)
        remember(level1, start)
    orphans = [item for item in expense if not item.get("fundingContractId") or not by_id.get(item.get("fundingContractId"))]
    orphan_free = [item for item in payments if not item.get("contractId") and (not item.get("fundingContractId") or not by_id.get(item.get("fundingContractId")))]
    if orphans or orphan_free:
        total = sum((excel_dec(item.get("amount")) for item in orphans), Decimal(0)) + sum((excel_dec(item.get("amount")) for item in orphan_free), Decimal(0))
        fact = sum((excel_dec(item.get("paidTotal")) for item in orphans), Decimal(0)) + sum((excel_dec(item.get("amount")) for item in orphan_free if item.get("status") == "paid"), Decimal(0))
        expense_limit += total
        expense_fact += fact
        write_excel_row(ws, analytics_values("Без источника финансирования", "", "", "", "", total.quantize(MONEY), fact.quantize(MONEY), fact_percent=excel_share(fact, total)), kinds, bold=True, fill=FILL_GROUP)
        start = ws.max_row + 1
        for item in orphans:
            write_contract(item, bold=True)
            write_specs(item, level2)
        write_free(orphan_free, level2)
        remember(level1, start)
    for start, end in level1:
        group_excel_rows(ws, start, end, 1)
    for start, end in level2:
        group_excel_rows(ws, start, end, 2)
    write_excel_row(ws, analytics_values("Итого доходы", "", "", "", "", income_limit.quantize(MONEY), income_fact.quantize(MONEY), fact_percent=excel_share(income_fact, income_limit)), kinds, bold=True, fill=FILL_TOTAL)
    write_excel_row(ws, analytics_values("Итого расходы", "", "", "", "", expense_limit.quantize(MONEY), expense_fact.quantize(MONEY), fact_percent=excel_share(expense_fact, expense_limit)), kinds, bold=True, fill=FILL_TOTAL)
    write_excel_row(ws, analytics_values("Остаток (доход − расход)", "", "", "", "", (income_limit - expense_limit).quantize(MONEY), (income_fact - expense_fact).quantize(MONEY)), kinds, bold=True, fill=FILL_TOTAL)
    others = [item for item in contracts if is_top(item) and item.get("otherExpense")]
    if others:
        ws.append([None] * len(headers))
        write_excel_row(ws, ["Прочие расходы"] + [None] * (len(headers) - 1), kinds, bold=True, fill=FILL_TOTAL)
        for item in others:
            write_contract(item, bold=True, fill=FILL_GROUP)
            start = ws.max_row + 1
            write_specs(item, [])
            group_excel_rows(ws, start, ws.max_row, 1)
    finish_excel_sheet(ws, [18, 42, 18, 28, 16, 22, 18, 18, 18, 24, 26, 18, 12])


def sheet_dictionaries(wb, dictionaries):
    ws = wb.active
    ws.title = "Справочники"
    kinds = "tt"
    ws.append(["Справочник", "Значение"])
    for kind in DICTIONARY_KINDS:
        items = [item for item in dictionaries if item.get("kind") == kind]
        items.sort(key=lambda item: text(item.get("value")).lower())
        write_excel_row(ws, [DICTIONARY_LABELS[kind], f"{len(items)}"], kinds, bold=True, fill=FILL_GROUP)
        if not items:
            continue
        start = ws.max_row + 1
        for item in items:
            write_excel_row(ws, [DICTIONARY_LABELS[kind], item.get("value") or ""], kinds)
        group_excel_rows(ws, start, ws.max_row, 1)
    finish_excel_sheet(ws, [24, 48])


def sheet_history(wb, history):
    ws = wb.active
    ws.title = "История"
    kinds = "dttttt"
    ws.append(["Дата", "Время", "Раздел", "Изменение", "Описание", "Пользователь"])
    grouped = {}
    for item in history:
        moment = item["ts"]
        grouped.setdefault(moment.date(), []).append(item)
    for day in sorted(grouped, reverse=True):
        items = grouped[day]
        write_excel_row(ws, [day, "", "", f"{len(items)}", "", ""], kinds, bold=True, fill=FILL_GROUP)
        start = ws.max_row + 1
        for item in items:
            moment = item["ts"]
            write_excel_row(ws, [
                moment.date(),
                moment.strftime("%H:%M:%S"),
                AUDIT_SECTION.get(item["entity_type"], "Система"),
                item["action"] or "",
                item["label"] or "",
                item["actor"] or "",
            ], kinds)
        group_excel_rows(ws, start, ws.max_row, 1)
    finish_excel_sheet(ws, [16, 14, 18, 32, 48, 20])


def build_excel(page: str, contracts, payments, dictionaries, history, addendums=()) -> bytes:
    wb = Workbook()
    if page == "contracts":
        sheet_contracts(wb, contracts)
    elif page == "addendums":
        sheet_addendums(wb, contracts, list(addendums))
    elif page == "payments":
        sheet_payments(wb, contracts, payments)
    elif page == "analytics":
        sheet_analytics(wb, contracts, payments)
    elif page == "dictionaries":
        sheet_dictionaries(wb, dictionaries)
    elif page == "history":
        sheet_history(wb, history)
    else:
        raise HTTPException(400, "Неизвестная страница")
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@app.get("/api/export")
def export_registry(request: Request):
    guard(request, ALL_GROUPS)
    with transaction() as conn:
        contracts, _, payments, _ = fetch_all(conn)
    content = export_csv(contracts, payments).encode("utf-8")
    return Response(content, media_type="text/csv; charset=utf-8", headers={"Content-Disposition": "attachment; filename=registry.csv"})


@app.get("/api/export.xlsx")
def export_excel(request: Request, page: str = Query("contracts")):
    if page not in EXCEL_PAGES:
        raise HTTPException(400, "Неизвестная страница")
    guard(request, AUDIT_GROUPS if page == "history" else ALL_GROUPS)
    with transaction() as conn:
        contracts, _, payments, addendums = fetch_all(conn)
        dictionaries = [dictionary_from_row(row) for row in conn.execute("SELECT * FROM dictionaries ORDER BY kind, lower(value)")]
        history = []
        if page == "history":
            history = conn.execute("SELECT ts, actor, action, entity_type, label FROM audit_log ORDER BY id DESC").fetchall()
    content = build_excel(page, contracts, payments, dictionaries, history, addendums)
    filenames = {
        "contracts": "contracts.xlsx",
        "addendums": "addendums.xlsx",
        "payments": "payments.xlsx",
        "analytics": "analytics.xlsx",
        "dictionaries": "dictionaries.xlsx",
        "history": "history.xlsx",
    }
    return Response(
        content,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filenames[page]}"'},
    )


@app.post("/api/import")
def import_registry(request: Request, payload: dict = Body()):
    actor = guard(request, EDIT_GROUPS)
    rows = csv_table(text(payload.get("text")))
    if not rows:
        return {"addedContracts": 0, "addedPayments": 0, "duplicates": [], "errors": ["В файле нет строк для загрузки"]}
    duplicates = []
    errors = []
    accepted_contracts = []
    accepted_payments = []
    with transaction() as conn:
        stored_contracts = [contract_from_row(row) for row in conn.execute("SELECT * FROM contracts")]
        stored_payments = conn.execute("SELECT contract_id, number FROM payments").fetchall()
        contracts_by_number = index_contracts(stored_contracts)
        payment_keys = {(row["contract_id"] or "", text(row["number"]).strip()) for row in stored_payments}
        contract_lines = {}
        payment_lines = {}
        contract_rows = []
        payment_rows = []
        for line, record in rows:
            section = fold(record["Раздел"])
            if section == "договор":
                contract_rows.append((line, record))
            elif section == "платеж":
                payment_rows.append((line, record))
            else:
                errors.append(RowProblem(line, "Раздел", "укажите «Договор» или «Платеж»").text())

        def reference_depth(item) -> int:
            record = item[1]
            if text(record.get("Договор-основание")).strip():
                return 2
            return 1 if text(record.get("Источник финансирования")).strip() else 0

        # сначала самостоятельные договоры, затем те, что ссылаются на источник, затем спецификации
        contract_rows.sort(key=reference_depth)
        for line, record in contract_rows:
            try:
                entity = contract_from_import(record, line, contracts_by_number)
            except RowProblem as problem:
                errors.append(problem.text())
                continue
            number = entity["number"]
            if number in contract_lines:
                duplicates.append(f"Строка {line}: дубль договора № {number}. Первая строка файла — {contract_lines[number]}.")
                continue
            if number in contracts_by_number:
                duplicates.append(f"Строка {line}: дубль договора № {number}. Такая запись уже есть в реестре.")
                continue
            contract_lines[number] = line
            contracts_by_number[number] = entity
            accepted_contracts.append(entity)
        for line, record in payment_rows:
            try:
                entity = payment_from_import(record, line, contracts_by_number)
            except RowProblem as problem:
                errors.append(problem.text())
                continue
            key = (entity["contractId"], entity["number"])
            marker = f"{entity['contractId']}|{entity['number']}"
            if not entity["number"]:
                # без номера дубль определить нельзя — принимаем строку как есть
                accepted_payments.append(entity)
                continue
            if marker in payment_lines:
                duplicates.append(f"Строка {line}: дубль платежа № {entity['number']}. Первая строка файла — {payment_lines[marker]}.")
                continue
            if key in payment_keys:
                duplicates.append(f"Строка {line}: дубль платежа № {entity['number']} по этому договору. Такая запись уже есть в реестре.")
                continue
            payment_lines[marker] = line
            payment_keys.add(key)
            accepted_payments.append(entity)
        def line_number(message: str) -> int:
            found = re.search(r"Строка (\d+)", message)
            return int(found.group(1)) if found else 0
        duplicates.sort(key=line_number)
        errors.sort(key=line_number)
        for entity in accepted_contracts:
            saved = apply_saved(conn, "contract", entity, 1)
            audit(conn, actor["username"], "Договор добавлен", "contract", saved["id"], f"№ {saved['number']} · {saved['counterparty']}", None, saved)
        for entity in accepted_payments:
            saved = apply_saved(conn, "payment", entity, 1)
            audit(conn, actor["username"], "Операция добавлена", "payment", saved["id"], f"{saved['number']} · {saved['amount']}", None, saved)
    return {
        "addedContracts": len(accepted_contracts),
        "addedPayments": len(accepted_payments),
        "duplicates": duplicates,
        "errors": errors,
    }


@app.get("/api/audit")
def list_audit(request: Request):
    guard(request, AUDIT_GROUPS)
    with transaction() as conn:
        rows = conn.execute(
            "SELECT id, ts, actor, action, entity_type, label FROM audit_log ORDER BY id DESC LIMIT 200"
        ).fetchall()
    return [{
        "id": str(row["id"]),
        "timestamp": row["ts"].isoformat(),
        "actor": row["actor"],
        "action": row["action"],
        "entityType": row["entity_type"],
        "label": row["label"],
    } for row in rows]


def create_entity(request: Request, kind: str, payload: dict):
    row = guard(request, EDIT_GROUPS)
    with transaction() as conn:
        if kind == "contract":
            entity = normalize_contract(payload, None)
            if load_contract(conn, entity["id"]):
                raise HTTPException(409, "Строку уже изменили")
            saved = apply_saved(conn, kind, entity, 1)
            label = f"№ {saved['number']} · {saved['counterparty']}"
            action = "Договор добавлен"
        elif kind == "specification":
            entity = normalize_spec(payload, None)
            if load_spec(conn, entity["id"]):
                raise HTTPException(409, "Строку уже изменили")
            saved = apply_saved(conn, kind, entity, 1)
            label = saved["number"]
            action = "Спецификация добавлена"
        elif kind == "dictionary":
            entity = normalize_dictionary(payload, None)
            if load_dictionary(conn, entity["id"]):
                raise HTTPException(409, "Строку уже изменили")
            if conn.execute("SELECT 1 FROM dictionaries WHERE kind = %s AND lower(value) = lower(%s)", (entity["kind"], entity["value"])).fetchone():
                raise HTTPException(400, "Такое значение уже есть в справочнике")
            saved = apply_saved(conn, kind, entity, 1)
            label = f"{DICTIONARY_LABELS[saved['kind']]} · {saved['value']}"
            action = "Значение справочника добавлено"
        elif kind == "addendum":
            entity = normalize_addendum(payload, None)
            if load_addendum(conn, entity["id"]):
                raise HTTPException(409, "Строку уже изменили")
            saved = apply_saved(conn, kind, entity, 1)
            label = f"ДС № {saved['number']} · {saved['amount']}"
            action = "Доп. соглашение добавлено"
        else:
            entity = normalize_payment(payload, None)
            if load_payment(conn, entity["id"]):
                raise HTTPException(409, "Строку уже изменили")
            saved = apply_saved(conn, kind, entity, 1)
            label = f"{saved['number']} · {saved['amount']}"
            action = "Операция добавлена"
        audit(conn, row["username"], action, kind if kind != "specification" else "specification", saved["id"], label, None, saved)
    return saved


def update_entity(request: Request, kind: str, entity_id: str, payload: dict):
    row = guard(request, EDIT_GROUPS)
    with transaction() as conn:
        current = LOADERS[kind](conn, entity_id)
        if not current:
            raise HTTPException(404, "Запись не найдена")
        if int(payload.get("version") or 0) != int(current["version"]):
            raise HTTPException(409, "Строку уже изменили")
        if kind == "contract":
            entity = normalize_contract({**payload, "id": entity_id}, current)
            action, label = "Договор изменен", f"№ {entity['number']} · {entity['counterparty']}"
        elif kind == "specification":
            entity = normalize_spec({**payload, "id": entity_id}, current)
            action, label = "Спецификация изменена", entity["number"]
        elif kind == "dictionary":
            entity = normalize_dictionary({**payload, "id": entity_id}, current)
            action, label = "Значение справочника изменено", f"{DICTIONARY_LABELS[entity['kind']]} · {entity['value']}"
        elif kind == "addendum":
            entity = normalize_addendum({**payload, "id": entity_id}, current)
            action, label = "Доп. соглашение изменено", f"ДС № {entity['number']} · {entity['amount']}"
        else:
            entity = normalize_payment({**payload, "id": entity_id}, current)
            action, label = "Операция изменена", f"{entity['number']} · {entity['amount']}"
        saved = apply_saved(conn, kind, entity, int(current["version"]) + 1)
        audit(conn, row["username"], action, kind, entity_id, label, current, public_copy(saved))
    return saved


def delete_entity(request: Request, kind: str, entity_id: str, version: int):
    row = guard(request, EDIT_GROUPS)
    with transaction() as conn:
        current = LOADERS[kind](conn, entity_id)
        if not current:
            raise HTTPException(404, "Запись не найдена")
        if version != int(current["version"]):
            raise HTTPException(409, "Строку уже изменили")
        if kind == "contract":
            before = contract_bundle(conn, entity_id)
            label = f"№ {current['number']} · {current['counterparty']}"
            action = "Договор удален"
            conn.execute("DELETE FROM contracts WHERE parent_contract_id = %s", (entity_id,))
            conn.execute("UPDATE contracts SET funding_contract_id = '' WHERE funding_contract_id = %s", (entity_id,))
            conn.execute("UPDATE payments SET funding_contract_id = '' WHERE funding_contract_id = %s", (entity_id,))
            conn.execute("DELETE FROM contracts WHERE id = %s AND version = %s", (entity_id, version))
        elif kind == "specification":
            before, label, action = current, current["number"], "Спецификация удалена"
            conn.execute("DELETE FROM specifications WHERE id = %s AND version = %s", (entity_id, version))
        elif kind == "dictionary":
            before, label, action = current, f"{DICTIONARY_LABELS[current['kind']]} · {current['value']}", "Значение справочника удалено"
            conn.execute("DELETE FROM dictionaries WHERE id = %s AND version = %s", (entity_id, version))
        elif kind == "addendum":
            before, label, action = current, f"ДС № {current['number']} · {current['amount']}", "Доп. соглашение удалено"
            conn.execute("DELETE FROM addendums WHERE id = %s AND version = %s", (entity_id, version))
        else:
            before, label, action = current, f"{current['number']} · {current['amount']}", "Операция удалена"
            conn.execute("DELETE FROM payments WHERE id = %s AND version = %s", (entity_id, version))
        audit(conn, row["username"], action, kind, entity_id, label, before, None)
    return JSONResponse({"ok": True})


@app.post("/api/contracts")
def post_contract(request: Request, payload: dict = Body()):
    return create_entity(request, "contract", payload)


@app.patch("/api/contracts/{entity_id}")
def patch_contract(request: Request, entity_id: str, payload: dict = Body()):
    return update_entity(request, "contract", entity_id, payload)


@app.delete("/api/contracts/{entity_id}")
def delete_contract(request: Request, entity_id: str, version: int):
    return delete_entity(request, "contract", entity_id, version)


@app.post("/api/specifications")
def post_spec(request: Request, payload: dict = Body()):
    return create_entity(request, "specification", payload)


@app.patch("/api/specifications/{entity_id}")
def patch_spec(request: Request, entity_id: str, payload: dict = Body()):
    return update_entity(request, "specification", entity_id, payload)


@app.delete("/api/specifications/{entity_id}")
def delete_spec(request: Request, entity_id: str, version: int):
    return delete_entity(request, "specification", entity_id, version)


@app.post("/api/payments")
def post_payment(request: Request, payload: dict = Body()):
    return create_entity(request, "payment", payload)


@app.patch("/api/payments/{entity_id}")
def patch_payment(request: Request, entity_id: str, payload: dict = Body()):
    return update_entity(request, "payment", entity_id, payload)


@app.delete("/api/payments/{entity_id}")
def delete_payment(request: Request, entity_id: str, version: int):
    return delete_entity(request, "payment", entity_id, version)


@app.post("/api/addendums")
def post_addendum(request: Request, payload: dict = Body()):
    return create_entity(request, "addendum", payload)


@app.patch("/api/addendums/{entity_id}")
def patch_addendum(request: Request, entity_id: str, payload: dict = Body()):
    return update_entity(request, "addendum", entity_id, payload)


@app.delete("/api/addendums/{entity_id}")
def delete_addendum(request: Request, entity_id: str, version: int):
    return delete_entity(request, "addendum", entity_id, version)


@app.post("/api/dictionaries")
def post_dictionary(request: Request, payload: dict = Body()):
    return create_entity(request, "dictionary", payload)


@app.patch("/api/dictionaries/{entity_id}")
def patch_dictionary(request: Request, entity_id: str, payload: dict = Body()):
    return update_entity(request, "dictionary", entity_id, payload)


@app.delete("/api/dictionaries/{entity_id}")
def delete_dictionary(request: Request, entity_id: str, version: int):
    return delete_entity(request, "dictionary", entity_id, version)


TABLES = {"contract": "contracts", "specification": "specifications", "payment": "payments", "dictionary": "dictionaries", "addendum": "addendums"}
LOADERS = {"contract": load_contract, "specification": load_spec, "payment": load_payment, "dictionary": load_dictionary, "addendum": load_addendum}


def restore_record(conn, kind: str, entity: dict):
    version = int(entity.get("version") or 1)
    row = conn.execute(f"SELECT version FROM {TABLES[kind]} WHERE id = %s", (entity["id"],)).fetchone()
    if row:
        version = int(row["version"]) + 1
    if kind == "contract":
        entity = {**entity, "parentContractId": entity.get("parentContractId") or "", "fundingContractId": entity.get("fundingContractId") or ""}
    if kind == "payment":
        entity = {**entity, "counterparty": entity.get("counterparty") or "", "fundingContractId": entity.get("fundingContractId") or ""}
    apply_saved(conn, kind, entity, version)


@app.post("/api/audit/{audit_id}/rollback")
def rollback(request: Request, audit_id: int):
    row = guard(request, ROLLBACK_GROUPS)
    with transaction() as conn:
        entry = conn.execute("SELECT * FROM audit_log WHERE id = %s", (audit_id,)).fetchone()
        if not entry:
            raise HTTPException(404, "Запись журнала не найдена")
        before = entry["before_state"]
        kind = entry["entity_type"]
        if isinstance(before, dict) and "entity" in before and "specifications" in before:
            restore_record(conn, "contract", before["entity"])
            for child in before.get("children") or []:
                restore_record(conn, "contract", child)
            for spec in before.get("specifications") or []:
                restore_record(conn, "specification", spec)
            for payment in before.get("payments") or []:
                restore_record(conn, "payment", payment)
            for addendum in before.get("addendums") or []:
                restore_record(conn, "addendum", addendum)
            after = before
        elif before is None:
            table = TABLES[kind]
            conn.execute(f"DELETE FROM {table} WHERE id = %s", (entry["entity_id"],))
            after = None
        else:
            restore_record(conn, kind, before)
            after = before
        audit(conn, row["username"], "Выполнен откат", kind, entry["entity_id"], f"Отменено: {entry['action']} — {entry['label']}", entry["after_state"], after)
    return {"ok": True}


@app.get("/")
def index_page():
    return FileResponse(ROOT / "static" / "index.html", headers={"Cache-Control": "no-cache"})
